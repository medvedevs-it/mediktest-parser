from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


TEST_HEADERS = ["№", "Вопрос", "Правильный ответ", "Ответ 2", "Ответ 3", "Ответ 4"]
EXERCISE_HEADERS = ["id", "header", "condition", "subject_id"]
QUESTION_HEADERS = ["id", "question", "single_answer", "exercise_id", "text_before", "text_after", "header", "order"]
ANSWER_HEADERS = ["id", "answer", "is_correct", "question_id"]


def rows_after_header(sheet):
    iterator = sheet.iter_rows(values_only=True)
    header = list(next(iterator, ()))
    return header, list(iterator)


def is_all_caps_text(value) -> bool:
    letters = [character for character in str(value or "") if character.isalpha()]
    return bool(letters) and all(not character.islower() for character in letters)


def has_markup(value) -> bool:
    return bool(re.search(r"<[^>]+>|&(?:nbsp|lt|gt|amp|quot);", str(value or ""), re.IGNORECASE))


def has_markdown_image(value) -> bool:
    return bool(re.search(r"!\[[^\]]*\]\(https?://[^)]+\)", str(value or "")))


def verify_tests(path: Path, expected: int | None, unified: bool = False) -> tuple[dict, list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    errors = []
    expected_sheets = (
        ["Tests", "exercise_rows", "questions", "answers"]
        if unified
        else ["Tests"]
    )
    if workbook.sheetnames != expected_sheets:
        errors.append(
            "Единая книга должна содержать Tests, exercise_rows, questions, answers."
            if unified
            else "Книга тестов должна содержать только лист Tests."
        )
        return {}, errors
    header, tests = rows_after_header(workbook["Tests"])
    if header != TEST_HEADERS:
        errors.append("Колонки Tests не соответствуют шаблону.")
    if expected is not None and len(tests) != expected:
        errors.append(f"Ожидалось {expected} тестов, получено {len(tests)}.")
    test_ids = [row[0] for row in tests]
    if len(set(test_ids)) != len(test_ids):
        errors.append("В книге тестов есть повторяющиеся номера.")
    missing_questions = sum(not str(row[1] or "").strip() for row in tests)
    missing_correct_answers = sum(not str(row[2] or "").strip() for row in tests)
    missing_incorrect_answers = sum(
        any(not str(value or "").strip() for value in row[3:6]) for row in tests
    )
    duplicate_answers = sum(
        len({str(value or "").strip() for value in row[2:6]}) != 4 for row in tests
    )
    all_caps_questions = sum(is_all_caps_text(row[1]) for row in tests)
    questions_with_markup = sum(has_markup(row[1]) for row in tests)
    checks = {
        "missing_questions": missing_questions,
        "missing_correct_answers": missing_correct_answers,
        "missing_incorrect_answers": missing_incorrect_answers,
        "duplicate_answers": duplicate_answers,
        "all_caps_questions": all_caps_questions,
        "questions_with_markup": questions_with_markup,
    }
    labels = {
        "missing_questions": "Пустых тестовых вопросов",
        "missing_correct_answers": "Тестов без правильного ответа",
        "missing_incorrect_answers": "Тестов без трёх неверных ответов",
        "duplicate_answers": "Тестов с повторяющимся текстом вариантов",
        "all_caps_questions": "Тестовых вопросов целиком в верхнем регистре",
        "questions_with_markup": "Тестовых вопросов с HTML-разметкой",
    }
    for key, count in checks.items():
        if count:
            errors.append(f"{labels[key]}: {count}.")
    return {
        "tests": len(tests),
        "unique_test_ids": len(set(test_ids)),
        "answers_per_test": 4,
        "correct_answer_column": "C",
        **checks,
    }, errors


def verify_cases(path: Path, expected: int | None, unified: bool = False) -> tuple[dict, list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    errors = []
    expected_sheets = (
        ["Tests", "exercise_rows", "questions", "answers"]
        if unified
        else ["exercise_rows", "questions", "answers"]
    )
    if workbook.sheetnames != expected_sheets:
        errors.append(
            "Единая книга должна содержать Tests, exercise_rows, questions, answers."
            if unified
            else "Книга ситуационных задач должна содержать только exercise_rows, questions, answers."
        )
        return {}, errors
    exercise_header, exercises = rows_after_header(workbook["exercise_rows"])
    question_header, questions = rows_after_header(workbook["questions"])
    answer_header, answers = rows_after_header(workbook["answers"])
    if exercise_header != EXERCISE_HEADERS:
        errors.append("Колонки exercise_rows не соответствуют шаблону.")
    if question_header != QUESTION_HEADERS:
        errors.append("Колонки questions не соответствуют шаблону.")
    if answer_header != ANSWER_HEADERS:
        errors.append("Колонки answers не соответствуют шаблону.")
    if expected is not None and len(exercises) != expected:
        errors.append(f"Ожидалось {expected} ситуационных задач, получено {len(exercises)}.")

    exercise_ids = [row[0] for row in exercises]
    question_ids = [row[0] for row in questions]
    answer_ids = [row[0] for row in answers]
    if len(set(exercise_ids)) != len(exercise_ids):
        errors.append("В exercise_rows есть повторяющиеся ID.")
    if len(set(question_ids)) != len(question_ids):
        errors.append("В questions есть повторяющиеся ID.")
    if len(set(answer_ids)) != len(answer_ids):
        errors.append("В answers есть повторяющиеся ID.")

    exercise_id_set = set(exercise_ids)
    question_id_set = set(question_ids)
    questions_by_exercise = Counter(row[3] for row in questions)
    options_by_question = Counter(row[3] for row in answers)
    correct_by_question = Counter(row[3] for row in answers if row[2] in (True, 1, "1", "true", "True"))
    empty_conditions = sum(not str(row[2] or "").strip() for row in exercises)
    empty_headers = sum(not str(row[1] or "").strip() for row in exercises)
    generic_headers = sum(
        bool(re.fullmatch(r"Ситуационная задача(?:\s+№?\s*\d+)?", str(row[1] or "").strip(), re.IGNORECASE))
        for row in exercises
    )
    conditions_with_html_markup = sum(has_markup(row[2]) for row in exercises)
    orphan_questions = sum(row[3] not in exercise_id_set for row in questions)
    orphan_answers = sum(row[3] not in question_id_set for row in answers)
    answers_by_question = {}
    for answer in answers:
        answers_by_question.setdefault(answer[3], []).append(answer)
    invalid_answer_ids = sum(
        str(answer[0]) != f"{question_id}{index}"
        for question_id, question_answers in answers_by_question.items()
        for index, answer in enumerate(question_answers, 1)
    )
    cases_without_questions = sum(questions_by_exercise[exercise_id] == 0 for exercise_id in exercise_ids)
    questions_without_options = sum(options_by_question[question_id] == 0 for question_id in question_ids)
    questions_without_keys = sum(correct_by_question[question_id] == 0 for question_id in question_ids)
    all_caps_case_questions = sum(is_all_caps_text(row[1]) for row in questions)
    case_questions_with_markup = sum(has_markup(row[1]) for row in questions)
    text_after_rows = sum(bool(str(row[5] or "").strip()) for row in questions)
    text_after_with_html_markup = sum(has_markup(row[5]) for row in questions)
    text_after_with_image_links = sum(has_markdown_image(row[5]) for row in questions)

    checks = {
        "empty_conditions": empty_conditions,
        "empty_headers": empty_headers,
        "generic_headers": generic_headers,
        "conditions_with_html_markup": conditions_with_html_markup,
        "orphan_questions": orphan_questions,
        "orphan_answers": orphan_answers,
        "invalid_answer_ids": invalid_answer_ids,
        "cases_without_questions": cases_without_questions,
        "questions_without_options": questions_without_options,
        "questions_without_keys": questions_without_keys,
        "all_caps_case_questions": all_caps_case_questions,
        "case_questions_with_markup": case_questions_with_markup,
        "text_after_with_html_markup": text_after_with_html_markup,
    }
    for label, count in checks.items():
        if count:
            errors.append(f"{label}: {count}.")
    return {
        "cases": len(exercises),
        "case_questions": len(questions),
        "case_answers": len(answers),
        "unique_case_ids": len(exercise_id_set),
        "text_after_rows": text_after_rows,
        "text_after_with_image_links": text_after_with_image_links,
        **checks,
    }, errors


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("tests_xlsx", type=Path)
    parser.add_argument("cases_xlsx", type=Path, nargs="?")
    parser.add_argument("--unified", action="store_true")
    parser.add_argument("--expected-tests", type=int)
    parser.add_argument("--expected-cases", type=int)
    args = parser.parse_args()
    if args.unified:
        cases_path = args.tests_xlsx
    elif args.cases_xlsx is None:
        parser.error("cases_xlsx is required unless --unified is used")
    else:
        cases_path = args.cases_xlsx
    tests, test_errors = verify_tests(args.tests_xlsx, args.expected_tests, unified=args.unified)
    cases, case_errors = verify_cases(cases_path, args.expected_cases, unified=args.unified)
    errors = test_errors + case_errors
    print(json.dumps({"status": "ok" if not errors else "error", "tests": tests, "cases": cases, "errors": errors}, ensure_ascii=False, indent=2))
    return 0 if not errors else 1


if __name__ == "__main__":
    raise SystemExit(main())
