import tempfile
import time
import json
import subprocess
import sys
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from medik_pilot.collectors.demo import DemoCollector
from medik_pilot.collectors.live import LiveSelftestCollector
from medik_pilot.domain import (
    CollectedItem,
    extract_case_diagnosis,
    normalize_question_text,
    normalize_text,
    payload_status,
)
from medik_pilot.exporter import (
    _case_text_after,
    _condition,
    _markdown_table,
    export_audit_xlsx,
    export_cases_xlsx,
    export_images_zip,
    export_json,
    export_tests_xlsx,
    export_xlsx,
)
from medik_pilot.images import (
    client_image_filename,
    detect_image_type,
    filename_for_image_content,
)
from medik_pilot.runner import RunManager
from medik_pilot.specialties import SUPPORTED_SPECIALTIES, normalize_specialty, package_title
from medik_pilot.storage import Storage, utc_now
from medik_pilot.app import RunRequest, app, catalog_counts, create_run, public_image
from fastapi import HTTPException
from pydantic import ValidationError


class DomainTests(unittest.TestCase):
    def test_pediatrics_is_a_supported_independent_specialty(self):
        self.assertEqual(
            SUPPORTED_SPECIALTIES,
            ("Лечебное дело", "Педиатрия"),
        )
        self.assertEqual(normalize_specialty("  педиатрия  "), "Педиатрия")
        self.assertEqual(package_title("Педиатрия"), "РЭ_Педиатрия (специалитет), 2026")
        with self.assertRaises(ValueError):
            normalize_specialty("Стоматология")

    def test_collectors_keep_selected_pediatrics_specialty(self):
        demo = DemoCollector(specialty="Педиатрия")
        self.assertEqual(
            demo.collect_attempt("test", 1)[0].payload["specialty"],
            "Педиатрия",
        )
        self.assertEqual(
            demo.collect_attempt("case", 1)[0].payload["specialty"],
            "Педиатрия",
        )
        live = LiveSelftestCollector(settings=None, specialty="Педиатрия")
        self.assertEqual(live.specialty, "Педиатрия")
        self.assertEqual(live.package_title, "РЭ_Педиатрия (специалитет), 2026")
        self.assertRegex("РЭ_Педиатрия (специалитет), 2026", live.package_pattern)
        self.assertRegex("Педиатрия, 2026", live.package_pattern)
        self.assertNotRegex("РЭ_Лечебное дело, 2026", live.package_pattern)

    def test_run_request_normalizes_and_validates_specialty(self):
        request = RunRequest(specialty=" педиатрия ")
        self.assertEqual(request.specialty, "Педиатрия")
        with self.assertRaises(ValidationError):
            RunRequest(specialty="Стоматология")

    def test_catalog_counts_endpoint_is_specialty_scoped(self):
        with patch("medik_pilot.app.storage") as mocked_storage:
            mocked_storage.export_bank_counts.return_value = {"test": 13, "case": 3}
            result = catalog_counts(" педиатрия ")
        self.assertEqual(result, {
            "specialty": "Педиатрия",
            "test": 13,
            "case": 3,
            "total": 16,
        })
        mocked_storage.export_bank_counts.assert_called_once_with("Педиатрия")
        with self.assertRaises(HTTPException):
            catalog_counts("Стоматология")

    def test_runner_and_storage_separate_pediatrics_bank(self):
        with tempfile.TemporaryDirectory() as directory:
            storage = Storage(Path(directory) / "specialties.db")
            manager = RunManager(storage)
            pediatric_config = {
                "source_mode": "demo",
                "material_type": "test",
                "specialty": "Педиатрия",
                "reference_tests": 1,
                "reference_cases": 0,
                "max_attempts": 1,
                "max_duration_minutes": 1,
                "delay_seconds": 0,
            }
            collector = manager._collector("demo", pediatric_config)
            self.assertEqual(collector.specialty, "Педиатрия")
            for run_id, specialty in (
                ("medical-run", "Лечебное дело"),
                ("pediatric-run", "Педиатрия"),
            ):
                config = dict(pediatric_config, specialty=specialty)
                storage.create_run(run_id, config)
                item = DemoCollector(specialty=specialty).collect_attempt("test", 1)[0]
                result, _version = storage.store_item(
                    run_id, "demo", specialty, item, 1
                )
                self.assertEqual(result, "new")
            self.assertEqual(
                storage.catalog_counts("demo", "Лечебное дело")["test"], 1
            )
            self.assertEqual(
                storage.catalog_counts("demo", "Педиатрия")["test"], 1
            )

    def test_client_archive_validator_accepts_clean_package_and_rejects_database(self):
        validator = Path(__file__).resolve().parent.parent / "scripts" / "verify-client-package.py"
        with tempfile.TemporaryDirectory() as directory:
            clean = Path(directory) / "clean.zip"
            with zipfile.ZipFile(clean, "w") as archive:
                archive.writestr(
                    "MedikTest-Collector/MedikTest-Collector.exe", b"binary"
                )
                archive.writestr("CLIENT_INSTRUCTIONS.md", "Инструкция")
                archive.writestr("VERSION.txt", "1.0.0")
            accepted = subprocess.run(
                [sys.executable, str(validator), str(clean)],
                capture_output=True,
                text=True,
                check=False,
            )
            self.assertEqual(accepted.returncode, 0, accepted.stderr)

            unsafe = Path(directory) / "unsafe.zip"
            with zipfile.ZipFile(unsafe, "w") as archive:
                archive.writestr(
                    "MedikTest-Collector/MedikTest-Collector.exe", b"binary"
                )
                archive.writestr("CLIENT_INSTRUCTIONS.md", "Инструкция")
                archive.writestr("VERSION.txt", "1.0.0")
                archive.writestr("MedikTest-Collector/_internal/data/pilot.db", b"db")
            rejected = subprocess.run(
                [sys.executable, str(validator), str(unsafe)],
                capture_output=True,
                text=True,
                check=False,
            )
            self.assertNotEqual(rejected.returncode, 0)
            self.assertIn("pilot.db", rejected.stderr)

    def test_normalized_hash_ignores_whitespace_and_position(self):
        first = CollectedItem("test", "1", {"question": "Текст   вопроса", "position": 1})
        second = CollectedItem("test", "1", {"question": " Текст вопроса ", "position": 99})
        self.assertEqual(first.content_hash, second.content_hash)

    def test_normalizes_markup_and_question_case(self):
        self.assertEqual(
            normalize_question_text("  концентрация 10<sup>9</sup>/л&nbsp;повышена? "),
            "Концентрация 10⁹/л повышена?",
        )
        self.assertEqual(
            normalize_question_text("ДЛЯ АОРТАЛЬНОГО СТЕНОЗА ХАРАКТЕРНО НАЛИЧИЕ"),
            "Для аортального стеноза характерно наличие",
        )
        self.assertEqual(
            normalize_question_text("Оцените уровень HbA1c и СКФ CKD-EPI"),
            "Оцените уровень HbA1c и СКФ CKD-EPI",
        )

    def test_case_revealed_sections_move_from_condition_to_text_after(self):
        payload = {
            "condition_sections": [
                {"title": "Ситуация", "text": "Ситуация Пациент обратился к врачу."},
                {"title": "Жалобы", "text": "Жалобы на слабость."},
                {"title": "Результаты лабораторных методов обследования", "text": "Результаты лабораторных методов обследования"},
                {"title": "Общий анализ крови", "text": "Общий анализ крови Лейкоциты 10<sup>9</sup>/л"},
                {"title": "Результаты инструментальных методов обследования", "text": "Результаты инструментальных методов обследования"},
                {"title": "ЭКГ", "text": "ЭКГ Ритм синусовый"},
                {"title": "Диагноз", "text": "Диагноз Артериальная гипертензия"},
            ],
            "questions": [
                {"question": "Какие лабораторные исследования необходимы?", "options": []},
                {"question": "Какие инструментальные исследования необходимы?", "options": []},
                {"question": "Какой основной диагноз?", "options": [{"text": "Артериальная гипертензия", "is_correct": True}]},
            ],
        }
        self.assertEqual(
            _condition(payload),
            "**Ситуация**\nПациент обратился к врачу.\n**Жалобы**\nна слабость.",
        )
        laboratory = _case_text_after(
            payload,
            {"question": "Какие лабораторные исследования необходимы?"},
            1,
        )
        self.assertIn("Общий анализ крови\nЛейкоциты 10⁹/л", laboratory)
        self.assertNotIn("ЭКГ", laboratory)
        diagnosis = _case_text_after(
            payload,
            {"question": "Какой основной диагноз?"},
            3,
        )
        self.assertEqual(diagnosis, "Диагноз\nАртериальная гипертензия")

    def test_case_diagnosis_prefers_confirmed_answer_for_header(self):
        payload = {
            "header": "Ситуационная задача 7",
            "condition_sections": [
                {"title": "Диагноз", "text": "Диагноз Возможный диагноз"},
            ],
            "questions": [
                {
                    "question": "На основании полученных данных поставьте диагноз",
                    "options": [
                        {"text": "АРТЕРИАЛЬНАЯ ГИПЕРТЕНЗИЯ", "is_correct": True},
                        {"text": "Острый бронхит", "is_correct": False},
                    ],
                }
            ],
        }
        self.assertEqual(
            extract_case_diagnosis(payload),
            "Артериальная гипертензия",
        )

    def test_case_diagnosis_ignores_questions_about_diagnostic_methods(self):
        payload = {
            "header": "Ситуационная задача 1",
            "condition_sections": [
                {"title": "Диагноз", "text": "Диагноз"},
                {
                    "title": "Гипертоническая болезнь III стадии. ХСН IIА стадии",
                    "text": "Гипертоническая болезнь III стадии. ХСН IIА стадии",
                },
                {
                    "title": "Гипертрофическая кардиомиопатия",
                    "text": "Гипертрофическая кардиомиопатия",
                },
            ],
            "questions": [
                {
                    "question": "С целью подтверждения диагноза ХСН следует определить",
                    "options": [
                        {"text": "Мозговой натрийуретический пептид", "is_correct": True},
                        {"text": "Глюкозу", "is_correct": False},
                    ],
                }
            ],
        }
        self.assertEqual(
            extract_case_diagnosis(payload),
            "Гипертоническая болезнь III стадии. ХСН IIА стадии",
        )

    def test_case_diagnosis_uses_strict_diagnosis_question_before_sections(self):
        payload = {
            "condition_sections": [
                {"title": "Диагноз", "text": "Диагноз Возможный диагноз"},
            ],
            "questions": [
                {
                    "question": "Предварительный диагноз пациента",
                    "options": [
                        {"text": "НЕСТАБИЛЬНАЯ СТЕНОКАРДИЯ", "is_correct": True},
                        {"text": "Стабильная стенокардия", "is_correct": False},
                    ],
                }
            ],
        }
        self.assertEqual(extract_case_diagnosis(payload), "Нестабильная стенокардия")

    def test_case_condition_converts_source_bold_html_to_markdown(self):
        payload = {
            "condition_sections": [
                {
                    "title": "Ситуация",
                    "text": "Ситуация Пациенту выполнен <b>общий анализ крови</b>.",
                },
            ],
            "questions": [],
        }
        condition = _condition(payload)
        self.assertIn("**Ситуация**", condition)
        self.assertIn("**общий анализ крови**", condition)
        self.assertNotIn("<b>", condition)

    def test_case_text_after_contains_downloaded_image_filename(self):
        digest = "a" * 64
        payload = {
            "condition_sections": [
                {"title": "Ситуация", "text": "Пациент обратился к врачу."},
                {
                    "title": "Результаты лабораторных методов обследования",
                    "text": "Общий анализ крови",
                    "images": [
                        {
                            "asset_kind": "image",
                            "filename": "table_101668019967.png",
                            "storage_name": digest + ".png",
                            "sha256": digest,
                            "alt": "Общий анализ крови",
                            "order": 1,
                            "downloaded": True,
                        }
                    ],
                },
            ],
            "questions": [
                {"question": "Какие лабораторные исследования необходимы?", "options": []},
            ],
        }
        text_after = _case_text_after(payload, payload["questions"][0], 1)
        self.assertIn("Общий анализ крови", text_after)
        self.assertIn(
            "![Общий анализ крови](http://127.0.0.1:8765/api/assets/images/{}/table_101668019967.png)".format(digest),
            text_after,
        )

    def test_case_text_after_converts_structured_table_to_markdown_in_source_order(self):
        digest = "b" * 64
        payload = {
            "condition_sections": [
                {"title": "Ситуация", "text": "Пациент обратился к врачу."},
                {
                    "title": "Результаты лабораторных методов обследования",
                    "text": "Общий анализ крови Гемоглобин 80,0",
                    "content_blocks": [
                        {"type": "text", "text": "Общий анализ крови"},
                        {
                            "type": "table",
                            "rows": [
                                ["Наименование", "Нормы", "Результат"],
                                ["Гемоглобин", "130,0 - 160,0", "80,0"],
                            ],
                        },
                        {"type": "image", "image_order": 1},
                    ],
                    "images": [
                        {
                            "asset_kind": "image",
                            "filename": "blood table.png",
                            "storage_name": digest + ".png",
                            "sha256": digest,
                            "alt": "Таблица анализов",
                            "order": 1,
                            "downloaded": True,
                        }
                    ],
                },
            ],
            "questions": [
                {"question": "Какие лабораторные исследования необходимы?", "options": []},
            ],
        }
        text_after = _case_text_after(payload, payload["questions"][0], 1)
        table = (
            "| Наименование | Нормы | Результат |\n"
            "| --- | --- | --- |\n"
            "| Гемоглобин | 130,0 - 160,0 | 80,0 |"
        )
        self.assertIn(table, text_after)
        self.assertLess(text_after.index(table), text_after.index("![Таблица анализов]"))
        self.assertIn("blood%20table.png", text_after)

    def test_structured_text_after_wins_over_flattened_question_fallback(self):
        payload = {
            "condition_sections": [
                {"title": "Ситуация", "text": "Пациент обратился к врачу."},
                {
                    "title": "Результаты лабораторных методов обследования",
                    "text": "Гемоглобин 80,0",
                    "content_blocks": [
                        {
                            "type": "table",
                            "rows": [
                                ["Наименование", "Нормы", "Результат"],
                                ["Гемоглобин", "130,0 - 160,0", "80,0"],
                            ],
                        }
                    ],
                },
            ],
            "questions": [
                {
                    "question": "Что обнаружено при обследовании?",
                    "text_after": (
                        "Результаты лабораторных методов обследования "
                        "Гемоглобин 130,0 - 160,0 80,0"
                    ),
                    "options": [],
                }
            ],
        }
        text_after = _case_text_after(payload, payload["questions"][0], 1)
        self.assertIn("| Наименование | Нормы | Результат |", text_after)
        self.assertNotEqual(text_after, payload["questions"][0]["text_after"])

    def test_unmatched_revealed_image_is_kept_in_question_text_after(self):
        digest = "f" * 64
        payload = {
            "condition_sections": [
                {"title": "Ситуация", "text": "Пациент обратился к врачу."},
                {
                    "title": "Дополнительная информация",
                    "text": "Снимок исследования",
                    "images": [
                        {
                            "asset_kind": "image",
                            "filename": "scan.png",
                            "storage_name": digest + ".png",
                            "sha256": digest,
                            "alt": "Снимок исследования",
                            "order": 1,
                            "downloaded": True,
                        }
                    ],
                },
            ],
            "questions": [
                {"question": "Выберите дальнейшую тактику", "options": []},
            ],
        }
        text_after = _case_text_after(payload, payload["questions"][0], 1)
        self.assertIn("![Снимок исследования]", text_after)
        self.assertIn("/api/assets/images/{}/scan.png".format(digest), text_after)

    def test_markdown_table_escapes_pipes_inside_cells(self):
        self.assertEqual(
            _markdown_table([["Показатель", "Норма"], ["Na | K", "135 | 5"]]),
            "| Показатель | Норма |\n| --- | --- |\n| Na \\| K | 135 \\| 5 |",
        )

    def test_markdown_table_names_blank_result_header(self):
        self.assertEqual(
            _markdown_table(
                [
                    ["Наименование", "Нормы", ""],
                    ["Гемоглобин", "130,0 - 160,0", "80,0"],
                ]
            ),
            (
                "| Наименование | Нормы | Результат |\n"
                "| --- | --- | --- |\n"
                "| Гемоглобин | 130,0 - 160,0 | 80,0 |"
            ),
        )

    def test_public_image_route_serves_only_content_addressed_image(self):
        digest = "c" * 64
        with tempfile.TemporaryDirectory() as directory:
            image_dir = Path(directory)
            target = image_dir / (digest + ".png")
            target.write_bytes(b"png")
            with patch("medik_pilot.app.IMAGE_DIR", image_dir):
                response = public_image(digest, "table.png")
            self.assertEqual(Path(response.path), target)
            self.assertEqual(response.media_type, "image/png")
            self.assertIn("immutable", response.headers["cache-control"])

    def test_public_image_route_corrects_legacy_png_name_for_jpeg_bytes(self):
        digest = "e" * 64
        with tempfile.TemporaryDirectory() as directory:
            image_dir = Path(directory)
            target = image_dir / (digest + ".png")
            target.write_bytes(b"\xff\xd8\xff\xe0jpeg-content")
            with patch("medik_pilot.app.IMAGE_DIR", image_dir):
                response = public_image(digest, "photo.jpg")
            self.assertEqual(Path(response.path), target)
            self.assertEqual(response.media_type, "image/jpeg")

    def test_image_signature_controls_filename_and_mime(self):
        content = b"\xff\xd8\xff\xe0jpeg-content"
        self.assertEqual(detect_image_type(content), (".jpg", "image/jpeg"))
        self.assertEqual(
            filename_for_image_content("001001208.png", content),
            ("001001208.jpg", "image/jpeg"),
        )

    def test_client_image_filename_preserves_source_basename(self):
        self.assertEqual(
            client_image_filename(
                "https://storage.example/cases/table_101668019967.png?token=secret"
            ),
            "table_101668019967.png",
        )

    def test_payload_status_requires_disclosed_key(self):
        hidden = {"question": "Вопрос", "options": [{"text": "A", "is_correct": None}, {"text": "B", "is_correct": None}]}
        ready = {"question": "Вопрос", "options": [{"text": "A", "is_correct": True}, {"text": "B", "is_correct": False}]}
        self.assertEqual(payload_status("test", hidden), "missing_key")
        self.assertEqual(payload_status("test", ready), "ready")

    def test_case_answer_count_distinguishes_single_and_multiple(self):
        self.assertEqual(
            LiveSelftestCollector._required_case_answer_count(
                "Выберите необходимые исследования", ["checkbox"] * 4, 4
            ),
            1,
        )
        self.assertEqual(
            LiveSelftestCollector._required_case_answer_count(
                "Выберите необходимые исследования (выберите 2)",
                ["checkbox"] * 4,
                4,
            ),
            2,
        )

    def test_mutating_collector_is_not_retried_after_failure(self):
        class FailingCollector:
            allow_answer_submission = True

            def __init__(self):
                self.calls = 0

            def collect_attempt(self, kind, attempt):
                self.calls += 1
                raise RuntimeError("external attempt changed")

        with tempfile.TemporaryDirectory() as directory:
            storage = Storage(Path(directory) / "retry.db")
            storage.create_run(
                "run-retry",
                {
                    "source_mode": "live",
                    "material_type": "test",
                    "specialty": "Лечебное дело",
                    "max_attempts": 1,
                    "max_duration_minutes": 1,
                    "delay_seconds": 0,
                },
            )
            manager = RunManager(storage)
            collector = FailingCollector()
            with self.assertRaisesRegex(RuntimeError, "одной безопасной попытки"):
                manager._collect_with_retry(collector, "run-retry", "test", 1)
            self.assertEqual(collector.calls, 1)

    def test_read_only_retry_resets_browser_session(self):
        class NoWaitEvent:
            def is_set(self):
                return False

            def wait(self, timeout):
                return False

        class FlakyCollector:
            allow_answer_submission = False

            def __init__(self):
                self.calls = 0
                self.resets = 0

            def collect_attempt(self, kind, attempt):
                self.calls += 1
                if self.calls == 1:
                    raise RuntimeError("temporary network failure")
                return []

            def reset_session(self):
                self.resets += 1

        with tempfile.TemporaryDirectory() as directory:
            storage = Storage(Path(directory) / "retry-readonly.db")
            storage.create_run(
                "run-retry-readonly",
                {
                    "source_mode": "live",
                    "material_type": "test",
                    "specialty": "Лечебное дело",
                    "max_attempts": 1,
                    "max_duration_minutes": 1,
                    "delay_seconds": 0,
                },
            )
            manager = RunManager(storage)
            manager._stop_event = NoWaitEvent()
            collector = FlakyCollector()
            self.assertEqual(
                manager._collect_with_retry(
                    collector, "run-retry-readonly", "test", 1
                ),
                [],
            )
            self.assertEqual(collector.calls, 2)
            self.assertEqual(collector.resets, 1)

    def test_safely_resumable_mutating_collector_is_retried(self):
        class NoWaitEvent:
            def is_set(self):
                return False

            def wait(self, timeout):
                return False

        class ResumableCollector:
            allow_answer_submission = True
            supports_safe_resume = True

            def __init__(self):
                self.calls = 0
                self.resets = 0

            def collect_attempt(self, kind, attempt):
                self.calls += 1
                if self.calls == 1:
                    raise RuntimeError("delayed XForms transition")
                return []

            def reset_session(self):
                self.resets += 1

        with tempfile.TemporaryDirectory() as directory:
            storage = Storage(Path(directory) / "retry-resumable.db")
            storage.create_run(
                "run-retry-resumable",
                {
                    "source_mode": "live",
                    "material_type": "test",
                    "specialty": "Лечебное дело",
                    "max_attempts": 1,
                    "max_duration_minutes": 1,
                    "delay_seconds": 0,
                },
            )
            manager = RunManager(storage)
            manager._stop_event = NoWaitEvent()
            collector = ResumableCollector()
            self.assertEqual(
                manager._collect_with_retry(
                    collector, "run-retry-resumable", "test", 1
                ),
                [],
            )
            self.assertEqual(collector.calls, 2)
            self.assertEqual(collector.resets, 1)

    def test_live_collector_waits_at_pause_checkpoint(self):
        calls = []
        collector = LiveSelftestCollector(
            settings=None,
            pause_callback=lambda: calls.append("pause"),
            stop_callback=lambda: False,
        )
        self.assertFalse(collector._checkpoint())
        self.assertEqual(calls, ["pause"])

    def test_live_collector_selects_visible_xforms_copy(self):
        class Candidate:
            def __init__(self, visible):
                self.visible = visible

            def is_visible(self):
                return self.visible

        class Locator:
            def __init__(self, candidates):
                self.candidates = candidates

            def count(self):
                return len(self.candidates)

            def nth(self, index):
                return self.candidates[index]

        collector = LiveSelftestCollector(settings=None)
        collector._page = type(
            "Page",
            (),
            {"wait_for_timeout": lambda self, milliseconds: None},
        )()
        hidden = Candidate(False)
        visible = Candidate(True)
        selected = collector._first_visible(
            [Locator([hidden, hidden]), Locator([hidden, visible])],
            timeout_ms=0,
        )
        self.assertIs(selected, visible)

    def test_live_collector_downloads_image_to_content_addressed_store(self):
        class Response:
            ok = True
            status = 200
            headers = {"content-type": "image/png"}

            def body(self):
                return b"png-image-content"

        class Request:
            def get(self, url, timeout):
                self.url = url
                self.timeout = timeout
                return Response()

        class Context:
            request = Request()

        class Page:
            url = "https://selftest.example/case/1"
            context = Context()

        with tempfile.TemporaryDirectory() as directory:
            collector = LiveSelftestCollector(settings=None, image_dir=Path(directory))
            collector._page = Page()
            asset = collector._download_image_asset(
                {"source_url": "/media/table_101668019967.png", "order": 1}
            )
            self.assertTrue(asset["downloaded"])
            self.assertEqual(asset["filename"], "table_101668019967.png")
            self.assertEqual(
                (Path(directory) / asset["storage_name"]).read_bytes(),
                b"png-image-content",
            )

    def test_review_row_uses_last_xforms_trigger(self):
        clicks = []

        class Trigger:
            def __init__(self, name):
                self.name = name

            def click(self, force=False):
                clicks.append((self.name, force))

        class Triggers:
            def __init__(self):
                self.items = [
                    Trigger("incorrect-status"),
                    Trigger("correct-status"),
                    Trigger("open-question"),
                ]

            def count(self):
                return len(self.items)

            @property
            def last(self):
                return self.items[-1]

        class Row:
            def locator(self, selector):
                self.selector = selector
                return Triggers()

        class Cell:
            def locator(self, selector):
                self.selector = selector
                return Row()

        class NumberCells:
            def count(self):
                return 81

            def nth(self, number):
                return Cell()

        class Page:
            def locator(self, selector):
                self.selector = selector
                return NumberCells()

        item = CollectedItem(
            "test",
            "question-2",
            {
                "question": "Второй вопрос",
                "question_number": 2,
                "options": [
                    {"text": "A", "is_correct": True},
                    {"text": "B", "is_correct": False},
                ],
            },
        )
        collector = LiveSelftestCollector(settings=None)
        collector._page = Page()
        collector._show_test_result_list = lambda: None
        collector._wait_for_test_review_question = lambda **kwargs: item
        self.assertIs(collector._open_test_review_question(2), item)
        self.assertEqual(clicks, [("open-question", True)])

    def test_live_attempt_rotation_uses_unprocessed_materials_only_when_mutating(self):
        hrefs = ["attempt-1", "attempt-2"]
        self.assertEqual(
            LiveSelftestCollector._select_test_href(
                hrefs, {"attempt-1"}, rotate_attempts=True
            ),
            "attempt-2",
        )
        self.assertIsNone(
            LiveSelftestCollector._select_test_href(
                hrefs, set(hrefs), rotate_attempts=True
            )
        )
        self.assertEqual(
            LiveSelftestCollector._select_test_href(
                hrefs, set(hrefs), rotate_attempts=False
            ),
            "attempt-1",
        )

        case_links = [("/case/1", "case-1"), ("/case/2", "case-2")]
        self.assertEqual(
            LiveSelftestCollector._select_case_links(
                case_links, {"case-1"}, True, 1
            ),
            [("/case/2", "case-2")],
        )
        self.assertEqual(
            LiveSelftestCollector._select_case_links(
                case_links, {"case-1", "case-2"}, True, 1
            ),
            [],
        )
        self.assertEqual(
            LiveSelftestCollector._select_case_links(
                case_links, {"case-1", "case-2"}, False, 1
            ),
            [("/case/1", "case-1")],
        )

    def test_demo_attempts_overlap(self):
        collector = DemoCollector()
        first = {item.source_id for item in collector.collect_attempt("test", 1)}
        second = {item.source_id for item in collector.collect_attempt("test", 2)}
        self.assertTrue(first & second)
        self.assertTrue(second - first)


class StorageTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.storage = Storage(Path(self.temp.name) / "test.db")
        self.config = {
            "source_mode": "demo", "material_type": "test", "specialty": "Лечебное дело",
            "max_attempts": 2, "max_duration_minutes": 10, "saturation_window": 3,
            "novelty_threshold": 0.02, "delay_seconds": 0,
        }
        self.storage.create_run("run-1", self.config)

    def tearDown(self):
        self.temp.cleanup()

    def test_new_duplicate_and_changed_version(self):
        first = CollectedItem("test", "source-1", {"question": "Первый текст"})
        changed = CollectedItem("test", "source-1", {"question": "Изменённый текст"})
        self.assertEqual(self.storage.store_item("run-1", "demo", "Лечебное дело", first, 1)[0], "new")
        first_seen_at = self.storage.export_rows("run-1")[0]["first_seen_at"]
        self.storage.update_run(
            "run-1",
            status="completed",
            finished_at=utc_now(),
        )
        time.sleep(0.002)
        second_config = dict(self.config)
        self.storage.create_run("run-2", second_config)
        self.assertEqual(self.storage.store_item("run-1", "demo", "Лечебное дело", first, 1)[0], "duplicate")
        self.assertEqual(self.storage.store_item("run-2", "demo", "Лечебное дело", changed, 2)[0], "changed")
        versions = self.storage.export_rows("run-2")
        self.assertEqual([row["version"] for row in versions], [2])
        self.assertEqual(versions[0]["first_seen_at"], first_seen_at)
        historical = self.storage.catalog_rows_for_run("run-1")
        self.assertEqual(historical[0]["version"], 1)
        self.assertEqual(historical[0]["payload"]["question"], "Первый текст")

    def test_deduplicates_same_content_with_different_source_ids(self):
        first = CollectedItem("case", "variant-1", {"condition": "Одинаковое условие"})
        second = CollectedItem("case", "variant-2", {"condition": "Одинаковое условие"})
        self.assertEqual(self.storage.store_item("run-1", "demo", "Лечебное дело", first, 1)[0], "new")
        self.assertEqual(self.storage.store_item("run-1", "demo", "Лечебное дело", second, 2)[0], "duplicate")
        self.assertEqual(len(self.storage.export_rows("run-1")), 1)

    def test_client_ids_are_stable(self):
        first = self.storage.client_entity_id("demo", "Лечебное дело", "case_question", "case-1:q1", 600000)
        repeated = self.storage.client_entity_id("demo", "Лечебное дело", "case_question", "case-1:q1", 600000)
        second = self.storage.client_entity_id("demo", "Лечебное дело", "case_question", "case-1:q2", 600000)
        self.assertEqual(first, repeated)
        self.assertEqual(second, first + 1)

    def test_bulk_client_ids_are_stable_and_continue_existing_sequence(self):
        first = self.storage.client_entity_id(
            "demo", "Лечебное дело", "case_answer", "answer-1", 50000
        )
        assigned = self.storage.client_entity_ids_bulk(
            "demo",
            "Лечебное дело",
            "case_answer",
            ["answer-1", "answer-2", "answer-3", "answer-2"],
            50000,
        )
        self.assertEqual(assigned["answer-1"], first)
        self.assertEqual(assigned["answer-2"], first + 1)
        self.assertEqual(assigned["answer-3"], first + 2)

    def test_full_reset_removes_material_bank_and_run_history(self):
        test_item = CollectedItem("test", "test-1", {"question": "Тестовый вопрос"})
        case_item = CollectedItem("case", "case-1", {"condition": "Условие задачи"})
        self.storage.store_item("run-1", "demo", "Лечебное дело", test_item, 1)
        self.storage.store_item("run-1", "demo", "Лечебное дело", case_item, 1)
        self.storage.stage_raw_item("run-1", test_item, 1)
        self.storage.add_event("run-1", "Событие перед сбросом")
        self.storage.client_entity_id(
            "demo", "Лечебное дело", "case_question", "case-1:q1", 600000
        )

        result = self.storage.reset_all_materials()

        self.assertEqual(result["deleted_tests"], 1)
        self.assertEqual(result["deleted_cases"], 1)
        self.assertEqual(result["deleted_runs"], 1)
        self.assertIsNone(self.storage.latest_run())
        self.assertEqual(self.storage.list_runs(), [])
        with self.storage.database() as db:
            for table in (
                "runs", "items", "run_items", "attempt_stats", "events",
                "client_entity_ids", "raw_captures",
            ):
                count = db.execute(
                    "SELECT COUNT(*) AS count FROM {}".format(table)
                ).fetchone()["count"]
                self.assertEqual(count, 0, table)

    def test_scoped_specialty_reset_preserves_other_bank_and_ids(self):
        pediatric_config = dict(self.config, specialty="Педиатрия")
        self.storage.create_run("pediatric-run", pediatric_config)
        medical_payload = {"question": "Медицинский вопрос", "options": [{"text": "Да", "is_correct": True}, {"text": "Нет", "is_correct": False}]}
        pediatric_payload = {"question": "Педиатрический вопрос", "options": [{"text": "Да", "is_correct": True}, {"text": "Нет", "is_correct": False}]}
        medical_item = CollectedItem("test", "same-source-id", medical_payload)
        pediatric_item = CollectedItem("test", "same-source-id", pediatric_payload)
        self.storage.store_item("run-1", "demo", "Лечебное дело", medical_item, 1)
        self.storage.store_item("pediatric-run", "demo", "Педиатрия", pediatric_item, 1)
        medical_id = self.storage.client_entity_id("demo", "Лечебное дело", "question", "same-source-id", 1)
        pediatric_id = self.storage.client_entity_id("demo", "Педиатрия", "question", "same-source-id", 1)
        self.assertEqual(medical_id, pediatric_id)

        result = self.storage.reset_specialty_materials("Педиатрия")

        self.assertEqual(result["deleted_tests"], 1)
        self.assertEqual(result["deleted_runs"], 1)
        self.assertEqual(self.storage.catalog_counts("demo", "Педиатрия"), {})
        self.assertEqual(self.storage.catalog_counts("demo", "Лечебное дело")["test"], 1)
        self.assertEqual(self.storage.list_runs()[0]["specialty"], "Лечебное дело")
        with self.storage.database() as db:
            self.assertEqual(
                db.execute(
                    "SELECT COUNT(*) AS count FROM client_entity_ids WHERE specialty = ?",
                    ("Педиатрия",),
                ).fetchone()["count"],
                0,
            )

    def test_manager_scoped_reset_creates_backup(self):
        pediatric_config = dict(self.config, specialty="Педиатрия")
        self.storage.create_run("pediatric-run", pediatric_config)
        item = CollectedItem(
            "test",
            "pediatric-only",
            {"question": "Педиатрический вопрос", "options": [{"text": "Да", "is_correct": True}, {"text": "Нет", "is_correct": False}]},
        )
        self.storage.store_item("pediatric-run", "demo", "Педиатрия", item, 1)
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with patch("medik_pilot.runner.DATA_DIR", root), patch("medik_pilot.runner.DB_PATH", self.storage.path), patch("medik_pilot.runner.EXPORT_DIR", root / "exports"), patch("medik_pilot.runner.IMAGE_DIR", root / "images"), patch("medik_pilot.runner.PROBE_DIR", root / "probes"):
                result = RunManager(self.storage).reset_specialty_materials("Педиатрия")
            self.assertTrue(Path(result["backup_path"]).is_file())
            self.assertEqual(self.storage.catalog_counts("demo", "Лечебное дело").get("test", 0), 0)
            self.assertEqual(self.storage.catalog_counts("demo", "Педиатрия"), {})

    def test_full_reset_is_blocked_while_run_thread_is_alive(self):
        class AliveThread:
            @staticmethod
            def is_alive():
                return True

        manager = RunManager(self.storage)
        manager._thread = AliveThread()
        with self.assertRaisesRegex(RuntimeError, "пока выполняется сбор"):
            manager.reset_all_materials()
        self.assertIsNotNone(self.storage.latest_run())

    def test_incomplete_material_is_kept_but_not_exported_to_excel(self):
        item = CollectedItem(
            "test",
            "missing-key",
            {
                "question": "Вопрос без раскрытого ключа",
                "options": [
                    {"text": "Первый", "is_correct": None},
                    {"text": "Второй", "is_correct": None},
                ],
            },
        )
        self.storage.store_item("run-1", "demo", "Лечебное дело", item, 1)
        json_path = export_json(self.storage, "run-1")
        xlsx_path = export_xlsx(self.storage, "run-1")
        document = json.loads(json_path.read_text(encoding="utf-8"))
        self.assertEqual(document["tests"], [])
        self.assertEqual(len(document["incomplete"]), 1)
        from openpyxl import load_workbook
        workbook = load_workbook(xlsx_path, read_only=True)
        self.assertEqual(workbook["Tests"].max_row, 1)

    def test_image_export_uses_filename_from_text_after_and_separate_zip(self):
        image_dir = Path(self.temp.name) / "images"
        image_dir.mkdir()
        digest = "d" * 64
        storage_name = digest + ".png"
        (image_dir / storage_name).write_bytes(b"image-bytes")
        item = CollectedItem(
            "case",
            "case-with-image",
            {
                "condition": "Пациент обратился к врачу.",
                "condition_sections": [
                    {"title": "Ситуация", "text": "Пациент обратился к врачу."},
                    {
                        "title": "Результаты лабораторных методов обследования",
                        "text": "Общий анализ крови",
                        "images": [
                            {
                                "asset_kind": "image",
                                "filename": "table_101668019967.png",
                                "storage_name": storage_name,
                                "sha256": digest,
                                "content_type": "image/png",
                                "size": 11,
                                "source_url": "https://example/table_101668019967.png",
                                "downloaded": True,
                            }
                        ],
                    },
                ],
                "questions": [
                    {
                        "question_number": 1,
                        "question": "Какие лабораторные исследования необходимы?",
                        "multiple": False,
                        "options": [
                            {"text": "Общий анализ крови", "is_correct": True},
                            {"text": "ЭКГ", "is_correct": False},
                        ],
                    }
                ],
                "expected_questions": 1,
            },
        )
        self.storage.store_item("run-1", "demo", "Лечебное дело", item, 1)
        xlsx_path = export_xlsx(self.storage, "run-1")
        images_path = export_images_zip(self.storage, "run-1", image_dir=image_dir)
        from openpyxl import load_workbook
        workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
        image_reference = workbook["questions"]["F2"].value
        self.assertIn("![", image_reference)
        self.assertIn("/api/assets/images/", image_reference)
        self.assertIn("table_101668019967.png", image_reference)
        with zipfile.ZipFile(images_path) as archive:
            self.assertEqual(
                archive.read("images/table_101668019967.png"),
                b"image-bytes",
            )
            manifest = json.loads(archive.read("image_manifest.json"))
            self.assertEqual(manifest["image_count"], 1)

    def test_case_export_writes_confirmed_diagnosis_to_exercise_header(self):
        item = CollectedItem(
            "case",
            "case-with-diagnosis",
            {
                "condition": "Пациент обратился к врачу.",
                "questions": [
                    {
                        "question_number": 1,
                        "question": "Какой диагноз следует установить?",
                        "multiple": False,
                        "options": [
                            {"text": "ПНЕВМОНИЯ", "is_correct": True},
                            {"text": "Бронхит", "is_correct": False},
                        ],
                    }
                ],
                "expected_questions": 1,
            },
        )
        self.storage.store_item("run-1", "demo", "Лечебное дело", item, 1)

        from openpyxl import load_workbook

        workbook = load_workbook(export_xlsx(self.storage, "run-1"), read_only=True)
        self.assertEqual(workbook["exercise_rows"]["B2"].value, "Пневмония")
        document = json.loads(export_json(self.storage, "run-1").read_text(encoding="utf-8"))
        self.assertEqual(document["cases"][0]["header"], "Пневмония")

    def test_interrupted_run_becomes_resumable(self):
        self.storage.update_run("run-1", status="running")
        self.storage.recover_stale_runs()
        run = self.storage.get_run("run-1")
        self.assertEqual(run["status"], "paused")
        self.assertEqual(run["stop_reason"], "application_restart")

    def test_error_count_is_available_for_run_and_history(self):
        self.storage.add_event("run-1", "Проверочная ошибка", "error")
        self.assertEqual(self.storage.get_run("run-1")["error_count"], 1)
        self.assertEqual(self.storage.list_runs()[0]["error_count"], 1)

    def test_raw_capture_is_persisted_before_attempt_finishes(self):
        item = CollectedItem(
            "test",
            "raw-question",
            {"question": "Нормализованный текст", "options": []},
            {"question_html": "<p>Исходный <b>текст</b></p>"},
        )
        self.storage.stage_raw_item("run-1", item, 1)
        self.assertEqual(self.storage.get_run("run-1")["raw_capture_count"], 1)


class LiveParserTests(unittest.TestCase):
    def test_extracts_case_question_before_first_option(self):
        body = "Далее\nКакое исследование выполнить первым?\nОбщий анализ крови\nБиохимия"
        self.assertEqual(
            LiveSelftestCollector._extract_case_question(body, "Общий анализ крови"),
            "Какое исследование выполнить первым?",
        )

    def test_required_case_answer_count_uses_instruction(self):
        self.assertEqual(
            LiveSelftestCollector._required_case_answer_count(
                "Пациенту следует рекомендовать (выберите 4)",
                ["checkbox"] * 6,
                6,
            ),
            4,
        )
        self.assertEqual(
            LiveSelftestCollector._required_case_answer_count(
                "Выберите один ответ",
                ["radio"] * 4,
                4,
            ),
            1,
        )


class PilotFlowTests(unittest.TestCase):
    def test_separate_client_export_routes_are_available(self):
        paths = {route.path for route in app.routes}
        self.assertIn("/api/runs/{run_id}/export.tests.xlsx", paths)
        self.assertIn("/api/runs/{run_id}/export.images.zip", paths)
        self.assertIn("/api/runs/{run_id}/export.cases.xlsx", paths)
        self.assertIn("/assets/mediktest-logo.png", paths)
        self.assertIn("/api/catalog/reset", paths)

    def test_transient_outer_cycle_error_does_not_fail_full_run(self):
        ready_item = CollectedItem(
            "test",
            "recovered-question",
            {
                "question": "Вопрос после восстановления",
                "options": [
                    {"text": "Верно", "is_correct": True},
                    {"text": "Неверно", "is_correct": False},
                ],
            },
        )

        class RecoveringManager(RunManager):
            def __init__(self, storage):
                super().__init__(storage)
                self.collection_calls = 0

            def _collector(self, mode, config, run_id=None):
                return DemoCollector()

            def _collect_with_retry(self, collector, run_id, kind, attempt):
                self.collection_calls += 1
                if self.collection_calls == 1:
                    raise RuntimeError("temporary result-page delay")
                return [ready_item]

        with tempfile.TemporaryDirectory() as directory:
            storage = Storage(Path(directory) / "outer-recovery.db")
            manager = RecoveringManager(storage)
            run_id = manager.start(
                {
                    "source_mode": "demo",
                    "material_type": "test",
                    "specialty": "Лечебное дело",
                    "reference_tests": 1,
                    "reference_cases": 0,
                    "verification_percent": 0,
                    "max_attempts": 2,
                    "max_requests": 10,
                    "max_duration_minutes": 10,
                    "delay_seconds": 0,
                    "error_retry_delay_seconds": 0,
                }
            )
            for _ in range(100):
                run = storage.get_run(run_id)
                if run and run["status"] in {
                    "completed", "partial", "failed", "stopped"
                }:
                    break
                time.sleep(0.01)
            self.assertEqual(run["status"], "completed")
            self.assertEqual(run["attempts_completed"], 2)
            self.assertEqual(run["unique_items"], 1)

    def test_repeated_empty_source_results_fail_with_clear_error(self):
        class EmptyCollector:
            def collect_attempt(self, kind, attempt):
                return []

        with tempfile.TemporaryDirectory() as directory:
            storage = Storage(Path(directory) / "empty-source.db")
            manager = RunManager(storage)
            config = {
                "source_mode": "demo", "material_type": "case", "specialty": "Лечебное дело",
                "reference_tests": 0, "reference_cases": 3, "verification_percent": 0,
                "max_attempts": 10, "max_requests": 100,
                "max_duration_minutes": 10, "delay_seconds": 0,
                "document_mode": "new", "document_name": "Пустой источник",
            }
            with patch.object(manager, "_collector", return_value=EmptyCollector()):
                run_id = manager.start(config)
                for _ in range(100):
                    run = storage.get_run(run_id)
                    if run and run["status"] in {"completed", "partial", "failed", "stopped"}:
                        break
                    time.sleep(0.01)
            self.assertEqual(run["status"], "failed")
            self.assertIn("не получены с сайта после 3 последовательных циклов", run["error_message"])
            self.assertEqual(run["requests_made"], 0)
            self.assertEqual(run["attempts_completed"], 2)
            self.assertEqual(run["unique_items"], 0)

    def test_reference_limits_cap_demo_run(self):
        with tempfile.TemporaryDirectory() as directory:
            storage = Storage(Path(directory) / "limited.db")
            manager = RunManager(storage)
            config = {
                "source_mode": "demo", "material_type": "both", "specialty": "Лечебное дело",
                "reference_tests": 1, "reference_cases": 1, "verification_percent": 0,
                "max_attempts": 3, "max_requests": 100,
                "max_duration_minutes": 10, "delay_seconds": 0,
            }
            run_id = manager.start(config)
            for _ in range(100):
                run = storage.get_run(run_id)
                if run and run["status"] in {"completed", "failed", "stopped"}:
                    break
                time.sleep(0.01)
            self.assertEqual(run["status"], "completed")
            self.assertEqual(run["items_seen"], 2)
            self.assertEqual(run["stop_reason"], "reference_verified")
            self.assertEqual(run["checkpoint"]["last_processed"]["attempt"], 1)
            self.assertIn(run["checkpoint"]["last_processed"]["kind"], {"test", "case"})

    def test_run_request_limits_and_zero_guard(self):
        request = RunRequest(max_tests=10, max_cases=10)
        self.assertEqual((request.reference_tests, request.reference_cases), (10, 10))
        tests_only = RunRequest(
            material_type="test",
            reference_tests=10,
            reference_cases=25,
        )
        self.assertEqual((tests_only.reference_tests, tests_only.reference_cases), (10, 0))
        new_document = RunRequest(
            document_mode="new",
            document_name="Контрольная выборка",
            reference_tests=5,
            reference_cases=2,
            verification_percent=15,
        )
        self.assertEqual(new_document.verification_percent, 0)
        self.assertTrue(new_document.allow_create_attempts)
        self.assertTrue(new_document.allow_answer_submission)
        with self.assertRaises(HTTPException):
            create_run(RunRequest(reference_tests=0, reference_cases=0))

    def test_demo_run_reaches_exports(self):
        with tempfile.TemporaryDirectory() as directory:
            storage = Storage(Path(directory) / "flow.db")
            manager = RunManager(storage)
            config = {
                "source_mode": "demo", "material_type": "both", "specialty": "Лечебное дело",
                "reference_tests": 10, "reference_cases": 3, "verification_percent": 15,
                "max_attempts": 5, "max_requests": 100,
                "max_duration_minutes": 10, "delay_seconds": 0,
            }
            run_id = manager.start(config)
            for _ in range(100):
                run = storage.get_run(run_id)
                if run and run["status"] in {"completed", "partial", "failed", "stopped"}:
                    break
                time.sleep(0.01)
            self.assertEqual(run["status"], "completed")
            self.assertGreaterEqual(run["items_seen"], 13)
            self.assertGreaterEqual(run["unique_items"], 13)
            self.assertEqual(run["stop_reason"], "reference_verified")
            json_path = export_json(storage, run_id)
            xlsx_path = export_xlsx(storage, run_id)
            tests_xlsx_path = export_tests_xlsx(storage, run_id)
            cases_xlsx_path = export_cases_xlsx(storage, run_id)
            self.assertTrue(json_path.exists())
            self.assertTrue(xlsx_path.exists())
            self.assertTrue(tests_xlsx_path.exists())
            self.assertTrue(cases_xlsx_path.exists())
            document = json.loads(json_path.read_text(encoding="utf-8"))
            self.assertGreaterEqual(len(document["cases"]), 3)
            self.assertTrue(document["cases"][0]["condition"])
            self.assertTrue(document["cases"][0]["questions"])
            self.assertNotIn("id", document["run"])
            self.assertIn("incomplete", document)
            from openpyxl import load_workbook
            workbook = load_workbook(xlsx_path, read_only=True)
            self.assertEqual(
                workbook.sheetnames,
                [
                    "Tests", "exercise_rows", "questions", "answers",
                ],
            )
            self.assertGreaterEqual(workbook["exercise_rows"].max_row, 4)
            self.assertGreaterEqual(workbook["questions"].max_row, 13)
            self.assertIs(workbook["questions"]["C2"].value, True)
            first_question_id = workbook["questions"]["A2"].value
            first_answer_id = workbook["answers"]["A2"].value
            self.assertEqual(str(first_answer_id), f"{first_question_id}1")
            tests_workbook = load_workbook(tests_xlsx_path, read_only=True)
            self.assertEqual(
                tests_workbook.sheetnames,
                ["Tests"],
            )
            tests_sheet = tests_workbook["Tests"]
            self.assertEqual(tests_sheet.max_row, 11)
            self.assertEqual(
                list(next(tests_sheet.iter_rows(values_only=True))),
                ["№", "Вопрос", "Правильный ответ", "Ответ 2", "Ответ 3", "Ответ 4"],
            )
            for test_row in tests_sheet.iter_rows(min_row=2, values_only=True):
                self.assertEqual(len(test_row), 6)
                self.assertTrue(all(str(value or "").strip() for value in test_row[1:]))
                self.assertEqual(len(set(test_row[2:6])), 4)
            cases_workbook = load_workbook(cases_xlsx_path, read_only=True)
            self.assertEqual(cases_workbook.sheetnames, ["exercise_rows", "questions", "answers"])
            self.assertGreaterEqual(cases_workbook["exercise_rows"].max_row, 4)

    def test_repeat_export_contains_full_catalog_and_separate_delta(self):
        with tempfile.TemporaryDirectory() as directory:
            storage = Storage(Path(directory) / "catalog-export.db")
            manager = RunManager(storage)
            config = {
                "source_mode": "demo", "material_type": "both", "specialty": "Лечебное дело",
                "reference_tests": 10, "reference_cases": 3, "verification_percent": 15,
                "max_attempts": 5, "max_requests": 100,
                "max_duration_minutes": 10, "delay_seconds": 0,
            }
            first_run_id = manager.start(config)
            for _ in range(100):
                first = storage.get_run(first_run_id)
                if first and first["status"] in {"completed", "partial", "failed", "stopped"}:
                    break
                time.sleep(0.01)
            self.assertEqual(first["status"], "completed")

            second_run_id = manager.start(config)
            for _ in range(100):
                second = storage.get_run(second_run_id)
                if second and second["status"] in {"completed", "partial", "failed", "stopped"}:
                    break
                time.sleep(0.01)
            self.assertEqual(second["status"], "completed")
            self.assertEqual(second["unique_items"], 0)
            self.assertGreaterEqual(second["items_seen"], 13)
            self.assertEqual(second["collected_by_kind"], {"case": 3, "test": 10})

            document = json.loads(export_json(storage, second_run_id).read_text(encoding="utf-8"))
            self.assertEqual(len(document["tests"]), 10)
            self.assertEqual(len(document["cases"]), 3)
            self.assertEqual(document["run_delta"], [])
            from openpyxl import load_workbook
            workbook = load_workbook(export_xlsx(storage, second_run_id), read_only=True)
            self.assertEqual(workbook["Tests"].max_row, 11)
            self.assertEqual(workbook["exercise_rows"].max_row, 4)

    def test_new_document_collects_requested_run_materials_from_source_even_when_known(self):
        with tempfile.TemporaryDirectory() as directory:
            storage = Storage(Path(directory) / "new-document.db")
            manager = RunManager(storage)
            base = {
                "source_mode": "demo", "material_type": "both", "specialty": "Лечебное дело",
                "reference_tests": 10, "reference_cases": 3, "verification_percent": 0,
                "max_attempts": 5, "max_requests": 100,
                "max_duration_minutes": 10, "delay_seconds": 0,
                "document_mode": "catalog", "document_name": "Общий банк",
            }
            catalog_run_id = manager.start(base)
            for _ in range(100):
                catalog_run = storage.get_run(catalog_run_id)
                if catalog_run and catalog_run["status"] in {"completed", "partial", "failed", "stopped"}:
                    break
                time.sleep(0.01)
            self.assertEqual(catalog_run["status"], "completed")

            sample = dict(base)
            sample.update({
                "document_mode": "new", "document_name": "Новая выборка",
                "reference_tests": 5, "reference_cases": 2,
            })
            sample_run_id = manager.start(sample)
            for _ in range(100):
                sample_run = storage.get_run(sample_run_id)
                if sample_run and sample_run["status"] in {"completed", "partial", "failed", "stopped"}:
                    break
                time.sleep(0.01)
            self.assertEqual(sample_run["status"], "completed")
            self.assertEqual(sample_run["collected_by_kind"], {"case": 2, "test": 5})
            self.assertEqual(sample_run["catalog_counts"], {"case": 3, "test": 10})
            self.assertEqual(sample_run["catalog_audit"]["test"]["confirmed"], 5)
            self.assertEqual(sample_run["catalog_audit"]["case"]["confirmed"], 2)
            self.assertEqual(sample_run["unique_items"], 0)
            self.assertEqual(sample_run["duplicate_items"], 7)
            self.assertEqual(sample_run["items_seen"], 7)
            self.assertEqual(sample_run["requests_made"], 7)
            self.assertEqual(sample_run["attempts_completed"], 1)

            from openpyxl import load_workbook
            workbook = load_workbook(export_xlsx(storage, sample_run_id), read_only=True)
            self.assertEqual(
                workbook.sheetnames,
                [
                    "Tests", "exercise_rows", "questions", "answers",
                ],
            )
            self.assertEqual(workbook["Tests"].max_row, 6)
            self.assertEqual(workbook["exercise_rows"].max_row, 3)
            self.assertEqual(workbook["Tests"]["A2"].value, 1)
            self.assertEqual(workbook["exercise_rows"]["A2"].value, 1)

    def test_new_document_collects_exactly_three_cases_from_source(self):
        with tempfile.TemporaryDirectory() as directory:
            storage = Storage(Path(directory) / "new-case-document.db")
            manager = RunManager(storage)
            base = {
                "source_mode": "demo", "material_type": "case", "specialty": "Лечебное дело",
                "reference_tests": 0, "reference_cases": 3, "verification_percent": 0,
                "max_attempts": 5, "max_requests": 100,
                "max_duration_minutes": 10, "delay_seconds": 0,
                "document_mode": "catalog", "document_name": "Общий банк",
            }
            catalog_run_id = manager.start(base)
            for _ in range(100):
                catalog_run = storage.get_run(catalog_run_id)
                if catalog_run and catalog_run["status"] in {"completed", "partial", "failed", "stopped"}:
                    break
                time.sleep(0.01)
            self.assertEqual(catalog_run["status"], "completed")

            sample = dict(base)
            sample.update({
                "document_mode": "new",
                "document_name": "Три ситуационные задачи",
            })
            sample_run_id = manager.start(sample)
            for _ in range(100):
                sample_run = storage.get_run(sample_run_id)
                if sample_run and sample_run["status"] in {"completed", "partial", "failed", "stopped"}:
                    break
                time.sleep(0.01)

            self.assertEqual(sample_run["status"], "completed")
            self.assertEqual(sample_run["collected_by_kind"], {"case": 3})
            # The demo source overlaps between attempts, so four source rows
            # are inspected across three attempts before three distinct cases
            # are reached.
            self.assertEqual(sample_run["items_seen"], 4)
            self.assertEqual(sample_run["duplicate_items"], 4)
            self.assertEqual(sample_run["requests_made"], 4)
            self.assertEqual(sample_run["attempts_completed"], 3)

            from openpyxl import load_workbook
            workbook = load_workbook(export_xlsx(storage, sample_run_id), read_only=True)
            self.assertEqual(workbook["Tests"].max_row, 1)
            self.assertEqual(workbook["exercise_rows"].max_row, 4)

    def test_catalog_update_progress_uses_current_run_but_export_keeps_full_bank(self):
        with tempfile.TemporaryDirectory() as directory:
            storage = Storage(Path(directory) / "catalog-update.db")
            manager = RunManager(storage)
            initial = {
                "source_mode": "demo", "material_type": "both", "specialty": "Лечебное дело",
                "reference_tests": 10, "reference_cases": 3, "verification_percent": 0,
                "max_attempts": 5, "max_requests": 100, "max_duration_minutes": 10,
                "delay_seconds": 0, "document_mode": "catalog", "document_name": "Общий банк",
            }
            first_id = manager.start(initial)
            for _ in range(100):
                first = storage.get_run(first_id)
                if first and first["status"] in {"completed", "partial", "failed", "stopped"}:
                    break
                time.sleep(0.01)
            self.assertEqual(first["status"], "completed")

            update = dict(initial)
            update.update({
                "material_type": "test", "reference_tests": 5, "reference_cases": 0,
                "document_name": "Общий банк",
            })
            update_id = manager.start(update)
            for _ in range(100):
                current = storage.get_run(update_id)
                if current and current["status"] in {"completed", "partial", "failed", "stopped"}:
                    break
                time.sleep(0.01)
            self.assertEqual(current["status"], "completed")
            self.assertEqual(current["collected_by_kind"].get("test"), 5)
            self.assertEqual(current["catalog_counts"], {"case": 3, "test": 10})

            from openpyxl import load_workbook
            workbook = load_workbook(export_xlsx(storage, update_id), read_only=True)
            self.assertEqual(workbook["Tests"].max_row, 11)
            self.assertEqual(workbook["exercise_rows"].max_row, 4)

            not_checked = storage.catalog_audit_items(
                update_id, audit_status="not_checked", limit=100
            )
            self.assertEqual(not_checked["total"], 8)
            self.assertFalse(not_checked["deletion_confirmed"])
            self.assertTrue(all(item["audit_status"] == "not_checked" for item in not_checked["items"]))
            self.assertTrue(all(item["times_seen"] >= 1 for item in not_checked["items"]))

            confirmed = storage.catalog_audit_items(
                update_id, kind="test", audit_status="confirmed", limit=100
            )
            self.assertEqual(confirmed["total"], 5)
            self.assertTrue(all(item["kind"] == "test" for item in confirmed["items"]))

            actuality = load_workbook(export_audit_xlsx(storage, update_id), read_only=True)
            self.assertEqual(actuality.sheetnames, ["Актуальность"])
            self.assertEqual(actuality["Актуальность"].max_row, 14)
            self.assertEqual(
                actuality["Актуальность"]["C2"].value,
                "Не проверен в запуске",
            )

    def test_hard_limit_marks_run_partial(self):
        with tempfile.TemporaryDirectory() as directory:
            storage = Storage(Path(directory) / "partial.db")
            manager = RunManager(storage)
            run_id = manager.start({
                "source_mode": "demo", "material_type": "test", "specialty": "Лечебное дело",
                "reference_tests": 16, "reference_cases": 0, "verification_percent": 0,
                "max_attempts": 1, "max_requests": 100, "max_duration_minutes": 10,
                "delay_seconds": 0,
            })
            for _ in range(100):
                run = storage.get_run(run_id)
                if run and run["status"] in {"completed", "partial", "failed", "stopped"}:
                    break
                time.sleep(0.01)
            self.assertEqual(run["status"], "partial")
            self.assertEqual(run["stop_reason"], "max_attempts")

    def test_repeats_do_not_stop_collection_before_hard_limit(self):
        class RepeatCollector:
            def collect_attempt(self, kind, attempt):
                return [
                    CollectedItem(
                        "test",
                        "repeat-only",
                        {
                            "question": "Один повторяющийся вопрос",
                            "options": [
                                {"text": "Верно", "is_correct": True},
                                {"text": "Неверно", "is_correct": False},
                            ],
                        },
                    )
                ]

            def close(self):
                return None

        class RepeatRunManager(RunManager):
            def _collector(self, mode, config, run_id=None):
                return RepeatCollector()

        with tempfile.TemporaryDirectory() as directory:
            storage = Storage(Path(directory) / "repeats.db")
            manager = RepeatRunManager(storage)
            run_id = manager.start(
                {
                    "source_mode": "demo",
                    "material_type": "test",
                    "specialty": "Лечебное дело",
                    "reference_tests": 2,
                    "reference_cases": 0,
                    "verification_percent": 0,
                    "max_attempts": 3,
                    "max_requests": 100,
                    "max_duration_minutes": 10,
                    "delay_seconds": 0,
                }
            )
            for _ in range(100):
                run = storage.get_run(run_id)
                if run and run["status"] in {"completed", "partial", "failed", "stopped"}:
                    break
                time.sleep(0.01)
            self.assertEqual(run["status"], "partial")
            self.assertEqual(run["stop_reason"], "max_attempts")
            self.assertEqual(run["attempts_completed"], 3)
            self.assertEqual(run["unique_items"], 1)
            self.assertEqual(run["duplicate_items"], 2)

    def test_request_limit_marks_run_partial(self):
        with tempfile.TemporaryDirectory() as directory:
            storage = Storage(Path(directory) / "request-limit.db")
            manager = RunManager(storage)
            run_id = manager.start({
                "source_mode": "demo", "material_type": "test", "specialty": "Лечебное дело",
                "reference_tests": 16, "reference_cases": 0, "verification_percent": 0,
                "max_attempts": 10, "max_requests": 1, "max_duration_minutes": 10,
                "delay_seconds": 0,
            })
            for _ in range(100):
                run = storage.get_run(run_id)
                if run and run["status"] in {"completed", "partial", "failed", "stopped"}:
                    break
                time.sleep(0.01)
            self.assertEqual(run["status"], "partial")
            self.assertEqual(run["stop_reason"], "max_requests")
            self.assertEqual(run["requests_made"], 1)

    def test_time_limit_marks_run_partial(self):
        with tempfile.TemporaryDirectory() as directory:
            storage = Storage(Path(directory) / "time-limit.db")
            manager = RunManager(storage)
            run_id = manager.start({
                "source_mode": "demo", "material_type": "test", "specialty": "Лечебное дело",
                "reference_tests": 16, "reference_cases": 0, "verification_percent": 0,
                "max_attempts": 10, "max_requests": 100, "max_duration_minutes": 0,
                "delay_seconds": 0,
            })
            for _ in range(100):
                run = storage.get_run(run_id)
                if run and run["status"] in {"completed", "partial", "failed", "stopped"}:
                    break
                time.sleep(0.01)
            self.assertEqual(run["status"], "partial")
            self.assertEqual(run["stop_reason"], "max_duration")

    def test_pause_and_resume_active_run(self):
        class SlowDemoCollector(DemoCollector):
            def collect_attempt(self, kind, attempt):
                time.sleep(0.1)
                return super().collect_attempt(kind, attempt)

        class SlowRunManager(RunManager):
            def _collector(self, mode, config, run_id=None):
                return SlowDemoCollector()

        with tempfile.TemporaryDirectory() as directory:
            storage = Storage(Path(directory) / "pause.db")
            manager = SlowRunManager(storage)
            run_id = manager.start({
                "source_mode": "demo", "material_type": "test", "specialty": "Лечебное дело",
                "reference_tests": 10, "reference_cases": 0, "verification_percent": 0,
                "max_attempts": 3, "max_requests": 100, "max_duration_minutes": 10,
                "delay_seconds": 0,
            })
            manager.pause(run_id)
            self.assertEqual(storage.get_run(run_id)["status"], "paused")
            manager.resume(run_id)
            for _ in range(200):
                run = storage.get_run(run_id)
                if run and run["status"] in {"completed", "partial", "failed", "stopped"}:
                    break
                time.sleep(0.01)
            self.assertEqual(run["status"], "completed")

    def test_manager_resumes_run_after_application_restart(self):
        with tempfile.TemporaryDirectory() as directory:
            storage = Storage(Path(directory) / "resume.db")
            config = {
                "source_mode": "demo", "material_type": "test", "specialty": "Лечебное дело",
                "reference_tests": 1, "reference_cases": 0, "verification_percent": 0,
                "max_attempts": 3, "max_requests": 100, "max_duration_minutes": 10,
                "allow_create_attempts": False, "allow_answer_submission": False,
                "delay_seconds": 0,
            }
            storage.create_run("resume-run", config)
            storage.update_run("resume-run", status="running")
            manager = RunManager(storage)
            self.assertEqual(storage.get_run("resume-run")["status"], "paused")
            manager.resume("resume-run")
            for _ in range(100):
                run = storage.get_run("resume-run")
                if run and run["status"] in {"completed", "partial", "failed", "stopped"}:
                    break
                time.sleep(0.01)
            self.assertEqual(run["status"], "completed")

    def test_resume_uses_saved_active_time_not_wall_clock_downtime(self):
        with tempfile.TemporaryDirectory() as directory:
            storage = Storage(Path(directory) / "active-time.db")
            config = {
                "source_mode": "demo", "material_type": "test", "specialty": "Лечебное дело",
                "reference_tests": 1, "reference_cases": 0, "verification_percent": 0,
                "max_attempts": 3, "max_requests": 100, "max_duration_minutes": 1,
                "allow_create_attempts": False, "allow_answer_submission": False,
                "delay_seconds": 0,
            }
            storage.create_run("old-paused-run", config)
            storage.update_run(
                "old-paused-run",
                status="running",
                started_at="2000-01-01T00:00:00+00:00",
                elapsed_seconds=0,
            )
            manager = RunManager(storage)
            manager.resume("old-paused-run")
            for _ in range(100):
                run = storage.get_run("old-paused-run")
                if run and run["status"] in {"completed", "partial", "failed", "stopped"}:
                    break
                time.sleep(0.01)
            self.assertEqual(run["status"], "completed")
            self.assertLess(run["elapsed_seconds"], 60)


if __name__ == "__main__":
    unittest.main()
