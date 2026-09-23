import io
import asyncio
import json
import tempfile
import time
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from medik_pilot.comparison import (ComparisonError, ComparisonManager, RETENTION,
    bank_values, compare_records, question_key, read_client, write_report)
from medik_pilot.domain import CollectedItem
from medik_pilot.storage import Storage


def payload(question):
    return {"question": question, "options": [
        {"text": "wrong1", "is_correct": False}, {"text": "right", "is_correct": True},
        {"text": "wrong2", "is_correct": False}, {"text": "wrong3", "is_correct": False}]}


class ComparisonTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def workbook(self, rows, headers=None):
        b = Workbook(); s = b.active; s.title = "Лист1"
        s.append(headers or ["index", "questiontext", "theme"])
        for row in rows:
            s.append(row)
        p = self.root / "client.xlsx"; b.save(p); b.close()
        return p

    def client(self, rows):
        return read_client(self.workbook(rows))["Лист1"]

    def test_normalization(self):
        self.assertEqual(question_key('<b>ВОПРОС</b>\u00a0«А» — Б'), question_key('вопрос "а" - б'))
        self.assertEqual(question_key('**Текст**<br> вопроса'), 'текст вопроса')
        self.assertEqual(question_key('10<sup>9</sup>/л'), question_key('10⁹/л'))
        self.assertEqual(question_key('CO<sub>2</sub>'), question_key('CO₂'))
        self.assertEqual(question_key('A &lt; 3 &gt; 1'), 'a < 3 > 1')

    def test_meaning_is_not_collapsed(self):
        for a,b in [('10⁹','109'),('не показано','показано'),('3 мг','5 мг'),('a < 3','a > 3'),('кто?','что?')]:
            self.assertNotEqual(question_key(a), question_key(b))

    def test_read_preserves_original_columns_and_rows(self):
        c = self.client([[17,'ВОПРОС','Тема'], [None,None,None],[18,'Вопрос','Другая']])
        self.assertEqual(c['headers'], ['index','questiontext','theme'])
        self.assertEqual([r['line'] for r in c['records']], [2,4])
        self.assertEqual(c['records'][1]['values'], [18,'Вопрос','Другая'])
        self.assertEqual(c['skipped'],1)

    def test_rejects_formulas_and_invalid_files(self):
        with self.assertRaisesRegex(ComparisonError,'формулы'):
            self.client([[1,'Вопрос','=1+1']])
        p=self.root/'invalid.xlsx';p.write_bytes(b'not a workbook')
        with self.assertRaises(ComparisonError): read_client(p)

    def test_requires_question_column_and_unique_headers(self):
        with self.assertRaisesRegex(ComparisonError,'questiontext'):
            read_client(self.workbook([], ['index','other']))
        with self.assertRaisesRegex(ComparisonError,'не повторяться'):
            read_client(self.workbook([], ['questiontext','questiontext']))

    def test_empty_sheet_is_previewable(self):
        self.assertEqual(self.client([])['records'], [])

    def test_limits(self):
        with patch('medik_pilot.comparison.MAX_ROWS', 1), self.assertRaises(ComparisonError):
            self.client([[1,'a','t'],[2,'b','t']])
        with patch('medik_pilot.comparison.MAX_COLS', 2), self.assertRaises(ComparisonError):
            self.client([[1,'a','t']])

    def test_duplicate_presence_and_reconciliation(self):
        c=self.client([[1,'A','t1'],[2,'a','t2'],[3,'B','t3'],[4,'b','t4']])
        bank=[{'source_id':i,'payload':payload(q)} for i,q in enumerate(['a','A','C'])]
        r=compare_records(c,bank);s=r['summary']
        self.assertEqual([x['values'][0] for x in r['missing_parser']], [3,4])
        self.assertEqual(len(r['missing_client']),1)
        self.assertEqual(s['matched_unique'],1)
        self.assertEqual(s['client_rows'],s['matched_client_rows']+s['missing_parser_rows'])
        self.assertEqual(s['bank_rows'],s['matched_bank_rows']+s['missing_client_rows'])
        self.assertEqual(s['client_unique'],s['matched_unique']+s['missing_parser_unique'])
        self.assertEqual(s['bank_unique'],s['matched_unique']+s['missing_client_unique'])

    def test_answers_do_not_affect_match(self):
        c=self.client([[1,'A','classification']])
        r=compare_records(c,[{'source_id':'1','payload':payload('a')}])
        self.assertEqual(r['summary']['matched_unique'],1)
        self.assertEqual(bank_values({'source_id':'1','payload':payload('a')})[2:], ['right','wrong1','wrong2','wrong3'])

    def test_permuted_answers_are_exported_once_but_conflicts_remain(self):
        import copy
        c=self.client([[1,'different','theme']])
        a={'source_id':'1','payload':payload('Question')}
        b=copy.deepcopy(a);b['source_id']='2';b['payload']['options'].reverse()
        conflict=copy.deepcopy(a);conflict['source_id']='3';conflict['payload']['options'][0]['text']='another'
        result=compare_records(c,[a,b,conflict])
        self.assertEqual(len(result['missing_client']),2)
        self.assertEqual(len(result['new_groups']['question']),3)
        self.assertEqual(result['summary']['exported_missing_client'],2)

    def test_current_bank_download_without_run_history_selection(self):
        from medik_pilot.exporter import export_current_bank_xlsx
        store=Storage(self.root/'bank.sqlite')
        for i,specialty in enumerate(['Лечебное дело','Педиатрия']):
            run=str(i);store.create_run(run,{'source_mode':'live','specialty':specialty,'material_type':'test',
                                            'max_attempts':1,'max_duration_minutes':1,'delay_seconds':0})
            store.store_item(run,'live',specialty,CollectedItem('test','same',payload(specialty)),1)
        self.assertEqual(store.export_bank_counts('Лечебное дело'),{'test':1,'case':0})
        self.assertEqual(store.export_bank_counts('Педиатрия'),{'test':1,'case':0})
        out=export_current_bank_xlsx(store,'Лечебное дело')
        try:
            b=load_workbook(out);self.assertEqual(b['Tests'].max_row,2)
            self.assertEqual(b['Tests']['B2'].value,'Лечебное дело');b.close()
        finally: out.unlink()

    def test_bank_counter_merges_permutations_and_preserves_answer_variants(self):
        import copy
        store=Storage(self.root/'counts.sqlite')
        store.create_run('r',{'source_mode':'live','specialty':'Педиатрия','material_type':'test',
                              'max_attempts':1,'max_duration_minutes':1,'delay_seconds':0})
        a=payload('Question');b=copy.deepcopy(a);b['options'].reverse()
        c=copy.deepcopy(a);c['options'][0]['text']='different'
        for i,p in enumerate([a,b,c]):
            store.store_item('r','live','Педиатрия',CollectedItem('test',str(i),p),1)
        self.assertEqual(store.export_bank_counts('Педиатрия')['test'],2)
        self.assertEqual(store.export_bank_counts('Лечебное дело')['test'],0)

    def test_run_limit_and_resume_keys_ignore_permuted_answers(self):
        import copy
        from medik_pilot.runner import RunManager
        from medik_pilot.collectors.demo import DemoCollector
        from medik_pilot.exporter import export_xlsx
        a=payload('First');b=copy.deepcopy(a);b['options'].reverse()
        samples=[a,b,payload('Second')]
        class Manager(RunManager):
            def _collector(self,*args,**kwargs): return DemoCollector()
            def _collect_with_retry(self,collector,run_id,kind,attempt):
                return [CollectedItem('test',str(attempt),samples[min(attempt-1,2)])]
        store=Storage(self.root/'limit.sqlite');manager=Manager(store)
        run_id=manager.start({'source_mode':'demo','specialty':'Лечебное дело','material_type':'test',
            'document_mode':'new','reference_tests':2,'reference_cases':0,'verification_percent':0,
            'max_attempts':3,'max_requests':10,'max_duration_minutes':10,'delay_seconds':0})
        for _ in range(500):
            run=store.get_run(run_id)
            if run['status'] in {'completed','failed','partial'}: break
            time.sleep(.01)
        self.assertEqual(run['status'],'completed')
        self.assertEqual(run['attempts_completed'],3)
        self.assertEqual(run['collected_by_kind']['test'],2)
        self.assertEqual(len(Storage(store.path).run_ready_identity_keys(run_id)['test']),2)
        with patch('medik_pilot.exporter.EXPORT_DIR',self.root):
            path=export_xlsx(store,run_id)
        b=load_workbook(path);self.assertEqual(b['Tests'].max_row,3);b.close()

    def test_report_preserves_text_and_duplicates(self):
        c=self.client([[1,'A','one'],[2,'a','two'],[3,'B','three']])
        c['records'][2]['values'][2]='=HYPERLINK("bad")'
        bank=[{'source_id':'=source','payload':payload('C')}]
        r=compare_records(c,bank);out=self.root/'report.xlsx'
        write_report(out,c,r,{'specialty':'Лечебное дело','captured_at':'2026-09-06','filename':'c.xlsx','sheet':'Лист1'})
        b=load_workbook(out)
        self.assertEqual(b.sheetnames,['Нет в базе клиента','Нет в банке парсера','Сводка','Дубли'])
        self.assertEqual(b.worksheets[1]['C4'].value,'=HYPERLINK("bad")')
        self.assertEqual(b.worksheets[1]['C4'].data_type,'s')
        self.assertEqual(b.worksheets[0]['A2'].data_type,'s')
        self.assertEqual(b['Дубли'].max_row,3)
        self.assertEqual(b['Дубли']['K2'].value,1)
        self.assertEqual(b['Дубли']['K3'].value,2)
        self.assertTrue(all(c.data_type!='f' for s in b for row in s for c in row))
        b.close()

    def test_snapshot_is_live_latest_scoped_and_readonly(self):
        store=Storage(self.root/'db.sqlite')
        for run,mode,specialty in [('med','live','Лечебное дело'),('ped','live','Педиатрия'),('demo','demo','Лечебное дело')]:
            store.create_run(run,{'source_mode':mode,'specialty':specialty,'material_type':'test',
                                 'max_attempts':10,'max_duration_minutes':10,'delay_seconds':0})
            store.store_item(run,mode,specialty,CollectedItem('test','same',payload(run)),1)
        store.store_item('med','live','Лечебное дело',CollectedItem('test','same',payload('latest')),2)
        store.store_item('med','live','Лечебное дело',CollectedItem('test','hidden',payload('old ready')),1)
        store.store_item('med','live','Лечебное дело',CollectedItem('test','hidden',{'question':'new incomplete'}),2)
        before=(self.root/'db.sqlite').read_bytes()
        snap=store.comparison_snapshot('Лечебное дело')
        self.assertEqual([x['payload']['question'] for x in snap['items']],['latest'])
        self.assertEqual(before,(self.root/'db.sqlite').read_bytes())
        self.assertEqual(store.comparison_snapshot('Педиатрия')['items'][0]['payload']['question'],'ped')
        with self.assertRaises(ValueError): store.comparison_snapshot('Unknown')

    def manager(self, bank=None):
        class FakeStorage:
            def comparison_snapshot(self,specialty):
                return {'captured_at':'2026-09-06T00:00:00+00:00','items':bank or []}
        m=ComparisonManager(FakeStorage(),self.root/'jobs')
        self.addCleanup(m.pool.shutdown,wait=True)
        return m

    def wait(self,m,identity):
        for _ in range(200):
            meta=m.get(identity)
            if meta['status'] not in {'reading','queued','comparing','exporting'}: return meta
            time.sleep(.01)
        self.fail('comparison did not finish')

    def test_manager_end_to_end(self):
        m=self.manager([{'source_id':'1','payload':payload('A')}])
        p=self.workbook([[1,'a',datetime(2026,9,6)]])
        j=m.upload(p.read_bytes(),'client.xlsx');j=self.wait(m,j['id'])
        self.assertEqual(j['status'],'ready');self.assertIsInstance(j['sheets'][0]['preview'][0][2],str)
        m.start(j['id'],'Лист1','Лечебное дело');j=self.wait(m,j['id'])
        self.assertEqual(j['status'],'completed');self.assertEqual(j['summary']['matched_unique'],1)
        self.assertTrue((m.directory(j['id'])/'report.xlsx').exists())

    def test_empty_bank_and_invalid_selection(self):
        m=self.manager();j=m.upload(self.workbook([[1,'a','t']]).read_bytes(),'client.xlsx');j=self.wait(m,j['id'])
        with self.assertRaises(ComparisonError): m.start(j['id'],'missing','Лечебное дело')
        with self.assertRaises(ValueError): m.start(j['id'],'Лист1','missing')
        m.start(j['id'],'Лист1','Педиатрия');j=self.wait(m,j['id'])
        self.assertEqual(j['status'],'failed');self.assertIn('пуст',j['error'])

    def test_expiry_path_traversal_and_busy(self):
        m=self.manager()
        with self.assertRaises(ComparisonError): m.get('../secret')
        with self.assertRaises(ComparisonError): m.upload(b'foo','bad.txt')
        identity='a'*32;d=m.directory(identity);d.mkdir()
        m.update(identity,created=time.time(),status='comparing')
        with self.assertRaises(ComparisonError): m.upload(b'foo','valid.xlsx')
        m.update(identity,created=time.time()-RETENTION-1,status='completed')
        with self.assertRaises(ComparisonError): m.get(identity)
        m.cleanup();self.assertFalse(d.exists())

    def test_restart_marks_interrupted(self):
        m=self.manager();identity='b'*32;m.directory(identity).mkdir()
        m.update(identity,created=time.time(),status='exporting')
        m2=ComparisonManager(m.storage,m.root);self.addCleanup(m2.pool.shutdown,wait=True)
        self.assertEqual(m2.get(identity)['status'],'interrupted')

    def test_api_upload_limits_and_errors(self):
        from medik_pilot.app import comparison_upload, comparison_status, comparison_report
        from fastapi import HTTPException
        class FakeRequest:
            async def stream(self):
                yield b'123'
                yield b'456'
        with patch('medik_pilot.app.MAX_BYTES',5), self.assertRaises(HTTPException) as ctx:
            asyncio.run(comparison_upload(FakeRequest(),'client.xlsx'))
        self.assertEqual(ctx.exception.status_code,413)
        m=self.manager()
        with patch('medik_pilot.app.comparisons',m):
            with self.assertRaises(HTTPException) as ctx: comparison_status('bad')
            self.assertEqual(ctx.exception.status_code,404)
            j=m.upload(self.workbook([[1,'a','t']]).read_bytes(),'client.xlsx');self.wait(m,j['id'])
            with self.assertRaises(HTTPException) as ctx: comparison_report(j['id'])
            self.assertEqual(ctx.exception.status_code,409)


if __name__=='__main__': unittest.main()
