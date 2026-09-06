import json
import re
import zipfile
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .config import EXPORT_DIR, IMAGE_DIR, ensure_directories, public_base_url
from .domain import extract_case_diagnosis, normalize_question_text, normalize_text
from .images import filename_for_image_content, iter_image_assets
from .storage import Storage


def export_item(row: Dict[str, Any]) -> Dict[str, Any]:
    item = {
        "source_id": row["source_id"],
        "version": row["version"],
        "content_hash": row["content_hash"],
        "status": row.get("status", "incomplete"),
        "first_seen_at": row["first_seen_at"],
        "last_seen_at": row["last_seen_at"],
        **_public_payload(row["payload"]),
    }
    if row["kind"] == "test":
        item.setdefault("single_answer", True)
    elif row["kind"] == "case":
        payload = row.get("payload") or {}
        item["header"] = extract_case_diagnosis(payload) or item.get("header")
        item["condition"] = _condition(payload)
        item["questions"] = [
            {
                **_public_payload(question),
                "text_after": _case_text_after(payload, question, order),
            }
            for order, question in enumerate(payload.get("questions") or [], 1)
        ]
    return item


def _public_payload(value: Any) -> Any:
    if isinstance(value, dict):
        return {
            key: _public_payload(nested)
            for key, nested in value.items()
            if key != "storage_name"
        }
    if isinstance(value, list):
        return [_public_payload(nested) for nested in value]
    return value


def _export_rows(storage: Storage, run_id: str) -> tuple[Dict[str, Any], List[Dict[str, Any]]]:
    run = storage.get_run(run_id)
    if not run:
        raise ValueError("Run not found")
    rows = (
        storage.export_rows(run_id)
        if run.get("document_mode", "catalog") == "new"
        else storage.catalog_rows_for_run(run_id)
    )
    return run, rows


def _image_manifest(rows: List[Dict[str, Any]], image_dir: Path = IMAGE_DIR) -> List[Dict[str, Any]]:
    by_filename: Dict[str, Dict[str, Any]] = {}
    case_ids: Dict[str, set[str]] = defaultdict(set)
    for row in rows:
        if row.get("kind") != "case" or row.get("status") != "ready":
            continue
        for asset in iter_image_assets(row.get("payload") or {}):
            if asset.get("downloaded") is not True:
                continue
            asset = _normalized_image_asset(asset, image_dir)
            filename = str(asset.get("filename") or "").strip()
            storage_name = str(asset.get("storage_name") or "").strip()
            digest = str(asset.get("sha256") or "").strip()
            if not filename or not storage_name:
                continue
            previous = by_filename.get(filename)
            if previous and previous.get("sha256") != digest:
                raise ValueError(
                    "Для имени изображения '{}' обнаружены разные файлы; экспорт остановлен, "
                    "чтобы не создать неверную привязку.".format(filename)
                )
            source_path = Path(image_dir) / storage_name
            if not source_path.is_file():
                continue
            by_filename[filename] = {
                "filename": filename,
                "storage_name": storage_name,
                "sha256": digest,
                "content_type": asset.get("content_type") or "",
                "size": int(asset.get("size") or source_path.stat().st_size),
                "source_url": asset.get("source_url") or "",
                "source_path": source_path,
            }
            case_ids[filename].add(str(row.get("source_id") or ""))
    result = []
    for filename in sorted(by_filename, key=str.casefold):
        item = by_filename[filename]
        item["case_source_ids"] = sorted(case_ids[filename])
        item["public_url"] = _public_image_url(item)
        result.append(item)
    return result


def export_json(storage: Storage, run_id: str) -> Path:
    ensure_directories()
    run, selected_rows = _export_rows(storage, run_id)
    catalog_rows = storage.catalog_rows_for_run(run_id)
    rows = selected_rows
    delta_rows = storage.export_rows(run_id)
    ready_rows = [row for row in rows if row.get("status") == "ready"]
    incomplete_rows = [row for row in rows if row.get("status") != "ready"]
    tests = [export_item(row) for row in ready_rows if row["kind"] == "test"]
    tests.sort(key=lambda item: (item.get("question_number") is None, item.get("question_number") or 0))
    cases = [export_item(row) for row in ready_rows if row["kind"] == "case"]
    incomplete = [export_item(row) for row in incomplete_rows]
    image_assets = _image_manifest(ready_rows)
    target = EXPORT_DIR / "{}.json".format(run_id)
    safe_run_fields = {
        "source_mode", "material_type", "specialty", "document_mode", "document_name",
        "status", "stop_reason",
        "created_at", "started_at", "finished_at", "reference_tests", "reference_cases",
        "verification_percent", "max_attempts", "max_requests", "max_duration_minutes",
        "attempts_completed", "items_seen", "unique_items", "duplicate_items",
        "changed_items", "requests_made",
    }
    target.write_text(
        json.dumps(
            {
                "run": {key: value for key, value in run.items() if key in safe_run_fields},
                "stats": run["attempt_stats"],
                "catalog_stats": {
                    "ready_tests": sum(
                        row["kind"] == "test" and row.get("status") == "ready"
                        for row in catalog_rows
                    ),
                    "ready_cases": sum(
                        row["kind"] == "case" and row.get("status") == "ready"
                        for row in catalog_rows
                    ),
                    "incomplete": sum(row.get("status") != "ready" for row in catalog_rows),
                },
                "export_stats": {
                    "scope": "current_run" if run.get("document_mode") == "new" else "catalog",
                    "ready_tests": len(tests),
                    "ready_cases": len(cases),
                    "incomplete": len(incomplete),
                },
                "catalog_audit": run.get("catalog_audit", {}),
                "run_delta": [
                    {
                        "kind": row["kind"],
                        "source_id": row["source_id"],
                        "version": row["version"],
                        "outcome": row["outcome"],
                        "status": row.get("status", "incomplete"),
                    }
                    for row in delta_rows
                    if row["outcome"] in {"new", "changed"}
                ],
                "tests": tests,
                "cases": cases,
                "image_assets": [
                    {
                        key: value
                        for key, value in asset.items()
                        if key not in {"source_path", "storage_name"}
                    }
                    for asset in image_assets
                ],
                "incomplete": incomplete,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    return target


_REVEALED_SECTION = re.compile(
    r"(?:^|\b)(?:результат|диагноз|дополнительн(?:ая|ые|ый)|получен\s+дополнительный)",
    re.IGNORECASE,
)


def _normalize_markdown_text(value: Any) -> str:
    """Normalize source HTML while preserving bold emphasis as Markdown."""
    prepared = re.sub(
        r"<(?:b|strong)[^>]*>(.*?)</(?:b|strong)>",
        lambda match: "**{}**".format(normalize_text(match.group(1))),
        str(value or ""),
        flags=re.IGNORECASE | re.DOTALL,
    )
    return normalize_text(prepared)


def _clean_section_text(title: str, text: str) -> str:
    title = normalize_text(title)
    plain_text = normalize_text(text)
    markdown_text = _normalize_markdown_text(text)
    if plain_text == title:
        return ""
    if title and plain_text.lower().startswith(title.lower() + " "):
        for prefix in ("**{}** ".format(title), title + " "):
            if markdown_text.lower().startswith(prefix.lower()):
                return markdown_text[len(prefix):].lstrip()
        return plain_text[len(title):].lstrip()
    return markdown_text


def _escape_markdown_cell(value: Any) -> str:
    return _normalize_markdown_text(value).replace("\\", "\\\\").replace("|", "\\|")


def _markdown_table(rows: List[List[Any]]) -> str:
    prepared = [
        [_escape_markdown_cell(cell) for cell in row]
        for row in rows
        if any(normalize_text(str(cell or "")) for cell in row)
    ]
    if not prepared:
        return ""
    width = max(len(row) for row in prepared)
    prepared = [row + [""] * (width - len(row)) for row in prepared]
    prepared[0] = [
        cell or ("Результат" if index == width - 1 else "Значение")
        for index, cell in enumerate(prepared[0])
    ]
    lines = ["| {} |".format(" | ".join(prepared[0]))]
    lines.append("| {} |".format(" | ".join(["---"] * width)))
    lines.extend("| {} |".format(" | ".join(row)) for row in prepared[1:])
    return "\n".join(lines)


def _public_image_url(asset: Dict[str, Any]) -> str:
    digest = str(asset.get("sha256") or "").strip().lower()
    filename = str(asset.get("filename") or "").strip()
    if not digest or not filename:
        return ""
    return "{}/api/assets/images/{}/{}".format(
        public_base_url(),
        quote(digest, safe=""),
        quote(filename, safe=""),
    )


def _normalized_image_asset(
    asset: Dict[str, Any], image_dir: Optional[Path] = None
) -> Dict[str, Any]:
    normalized = dict(asset)
    storage_name = str(normalized.get("storage_name") or "").strip()
    if not storage_name:
        return normalized
    source_path = Path(image_dir or IMAGE_DIR) / storage_name
    if not source_path.is_file():
        return normalized
    try:
        filename, content_type = filename_for_image_content(
            str(normalized.get("filename") or "image.png"),
            source_path.read_bytes(),
        )
    except OSError:
        return normalized
    if content_type:
        normalized["filename"] = filename
        normalized["content_type"] = content_type
    return normalized


def _markdown_image(asset: Dict[str, Any]) -> str:
    if asset.get("downloaded") is not True:
        return ""
    asset = _normalized_image_asset(asset)
    url = _public_image_url(asset)
    if not url:
        return ""
    filename = str(asset.get("filename") or "").strip()
    alt = normalize_text(str(asset.get("alt") or "")) or Path(filename).stem or filename
    alt = alt.replace("\\", "\\\\").replace("[", "\\[").replace("]", "\\]")
    return "![{}]({})".format(alt, url)


def _section_markdown(section: Dict[str, Any]) -> str:
    title = normalize_text(str(section.get("title") or ""))
    assets = {
        int(asset.get("order") or index): asset
        for index, asset in enumerate(section.get("images") or [], 1)
    }
    used_image_orders = set()
    blocks: List[str] = []
    content_blocks = section.get("content_blocks") or []
    for block in content_blocks:
        block_type = str(block.get("type") or "")
        if block_type == "text":
            text = _normalize_markdown_text(block.get("text"))
            if not blocks:
                text = _clean_section_text(title, text)
            if text:
                blocks.append(text)
        elif block_type == "table":
            table = _markdown_table(list(block.get("rows") or []))
            if table:
                blocks.append(table)
        elif block_type == "image":
            order = int(block.get("image_order") or 0)
            image = _markdown_image(assets.get(order) or {})
            if image:
                blocks.append(image)
                used_image_orders.add(order)
    if not content_blocks:
        text = _clean_section_text(title, str(section.get("text") or ""))
        if text:
            blocks.append(text)
    for order, asset in sorted(assets.items()):
        if order in used_image_orders:
            continue
        image = _markdown_image(asset)
        if image:
            blocks.append(image)
    return "\n\n".join(blocks)


def _split_case_sections(payload: Dict[str, Any]) -> tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    condition_sections: List[Dict[str, Any]] = []
    revealed_sections: List[Dict[str, Any]] = []
    revealed = False
    for section in payload.get("condition_sections") or []:
        title = normalize_text(str(section.get("title") or ""))
        text = _clean_section_text(title, str(section.get("text") or ""))
        if _REVEALED_SECTION.search(title):
            revealed = True
        target = revealed_sections if revealed else condition_sections
        if title or text or section.get("content_blocks") or section.get("images"):
            target.append({
                "title": title,
                "text": text,
                "content_blocks": list(section.get("content_blocks") or []),
                "images": list(section.get("images") or []),
            })
    return condition_sections, revealed_sections


def _format_sections(sections: List[Dict[str, Any]], emphasize_titles: bool = True) -> str:
    blocks = []
    for section in sections:
        title = section["title"]
        text = _section_markdown(section)
        if emphasize_titles and title:
            blocks.append("**{}**{}".format(title, "\n" + text if text else ""))
        elif title:
            blocks.append("{}{}".format(title, "\n" + text if text else ""))
        elif text:
            blocks.append(text)
    return "\n".join(blocks)


def _revealed_groups(sections: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    groups: List[Dict[str, Any]] = []
    for section in sections:
        title = section["title"]
        if _REVEALED_SECTION.search(title) or not groups:
            category = (
                "laboratory" if re.search(r"лабораторн", title, re.IGNORECASE)
                else "instrumental" if re.search(r"инструментальн|рентгенолог", title, re.IGNORECASE)
                else "diagnosis" if re.search(r"диагноз", title, re.IGNORECASE)
                else "general"
            )
            groups.append({"category": category, "sections": []})
        groups[-1]["sections"].append(section)
    return groups


def _text_overlap_score(left: str, right: str) -> int:
    left_tokens = set(re.findall(r"[0-9a-zа-яё]+", left.lower(), re.IGNORECASE))
    right_tokens = set(re.findall(r"[0-9a-zа-яё]+", right.lower(), re.IGNORECASE))
    if not left_tokens or not right_tokens:
        return 0
    ratio = len(left_tokens & right_tokens) / min(len(left_tokens), len(right_tokens))
    return int(ratio * 300) if ratio >= 0.45 else 0


def _case_export_structure(payload: Dict[str, Any]) -> tuple[List[Dict[str, str]], Dict[int, str]]:
    condition_sections, revealed_sections = _split_case_sections(payload)
    assignments: Dict[int, List[Dict[str, Any]]] = {}
    unassigned: List[Dict[str, Any]] = []
    questions = payload.get("questions") or []
    keyword_by_category = {
        "laboratory": r"лабораторн",
        "instrumental": r"инструментальн|рентгенолог|узи|экг|обследован",
        "diagnosis": r"диагноз",
        "general": r"обследован|исследован|результат|дополнительн",
    }
    for group in _revealed_groups(revealed_sections):
        group_text = normalize_text(_format_sections(group["sections"], emphasize_titles=False)).lower()
        candidates = []
        for order, question in enumerate(questions, 1):
            question_text = normalize_question_text(str(question.get("question") or ""))
            explicit_text_after = normalize_text(
                str(question.get("text_after") or "")
            ).lower()
            score = 0
            score += _text_overlap_score(group_text, explicit_text_after)
            if re.search(keyword_by_category[group["category"]], question_text, re.IGNORECASE):
                score += 100
            for option in question.get("options") or []:
                if option.get("is_correct") is not True:
                    continue
                option_text = normalize_text(str(option.get("text") or "")).lower()
                if len(option_text) >= 5 and option_text in group_text:
                    score += min(len(option_text), 80)
            if score:
                candidates.append((score, -order, order))
        if candidates:
            order = max(candidates)[2]
            assignments.setdefault(order, []).extend(group["sections"])
        elif questions:
            fallback_order = (
                len(questions)
                if group["category"] == "diagnosis"
                else min(2, len(questions))
                if group["category"] == "instrumental"
                else 1
            )
            assignments.setdefault(fallback_order, []).extend(group["sections"])
        else:
            unassigned.extend(group["sections"])
    return condition_sections + unassigned, {
        order: _format_sections(sections, emphasize_titles=False)
        for order, sections in assignments.items()
    }


def _condition(payload: Dict[str, Any]) -> str:
    sections = payload.get("condition_sections") or []
    if sections:
        condition_sections, _text_after = _case_export_structure(payload)
        return _format_sections(condition_sections)
    return _normalize_markdown_text(payload.get("condition"))


def _case_text_after(payload: Dict[str, Any], question: Dict[str, Any], order: int) -> Optional[str]:
    explicit = _normalize_markdown_text(question.get("text_after"))
    explicit_images = [
        image
        for image in (_markdown_image(asset) for asset in iter_image_assets(question))
        if image
    ]
    _condition_sections, text_after = _case_export_structure(payload)
    assigned = text_after.get(order) or ""
    # Structured sections preserve real HTML tables and image order. The
    # question-level text_after is only the flattened fallback from the page.
    combined = []
    for item in [assigned or explicit, *explicit_images]:
        if item and item not in combined:
            combined.append(item)
    return "\n\n".join(combined) or None


def _ordered_test_options(payload: Dict[str, Any]) -> List[tuple[int, Dict[str, Any]]]:
    """Return the client's required order: one correct answer, then three incorrect."""
    indexed = list(enumerate(payload.get("options") or [], 1))
    correct = [(index, option) for index, option in indexed if option.get("is_correct") is True]
    incorrect = [(index, option) for index, option in indexed if option.get("is_correct") is False]
    if len(correct) != 1 or len(incorrect) != 3 or len(indexed) != 4:
        raise ValueError(
            "Тест '{}' не соответствует шаблону: нужен один правильный и три "
            "неправильных ответа.".format(payload.get("question") or "")
        )
    return correct + incorrect


def _style_tests_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E79")
    blue_fill = PatternFill("solid", fgColor="DEEBF7")
    green_fill = PatternFill("solid", fgColor="E2EFDA")
    orange_fill = PatternFill("solid", fgColor="FCE4D6")
    white_font = Font(name="Carlito", size=11, bold=True, color="FFFFFF")
    body_font = Font(name="Carlito", size=11)
    thin = Side(style="thin", color="000000")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in range(2, sheet.max_row + 1):
        for column in range(1, 7):
            cell = sheet.cell(row, column)
            cell.font = body_font
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if column == 1:
                cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
            elif column == 2:
                cell.fill = blue_fill
            elif column == 3:
                cell.fill = green_fill
            else:
                cell.fill = orange_fill
        sheet.row_dimensions[row].height = 45
    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {
        "A": 10,
        "B": 60,
        "C": 48,
        "D": 48,
        "E": 48,
        "F": 48,
    }.items():
        sheet.column_dimensions[column].width = width


def _style_template_case_sheet(sheet, widths: Dict[str, float]) -> None:
    thin = Side(style="thin", color="D9D9D9")
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows():
        for cell in row:
            cell.font = Font(name="Carlito", size=11)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row].height = 45


def _ready_rows(storage: Storage, run_id: str) -> tuple[Dict[str, Any], List[Dict[str, Any]], List[Dict[str, Any]]]:
    run = storage.get_run(run_id)
    if not run:
        raise ValueError("Run not found")
    rows = (
        storage.export_rows(run_id)
        if run.get("document_mode", "catalog") == "new"
        else storage.catalog_rows_for_run(run_id)
    )
    tests_rows = [row for row in rows if row["kind"] == "test" and row.get("status") == "ready"]
    case_rows = [row for row in rows if row["kind"] == "case" and row.get("status") == "ready"]
    return run, tests_rows, case_rows


def _add_tests_sheet(
    workbook: Workbook,
    storage: Storage,
    tests_rows: List[Dict[str, Any]],
    stable_ids: bool = True,
) -> None:
    tests_sheet = workbook.active
    tests_sheet.title = "Tests"
    tests_sheet.append([
        "№",
        "Вопрос",
        "Правильный ответ",
        "Ответ 2",
        "Ответ 3",
        "Ответ 4",
    ])
    source_mode = tests_rows[0]["source_mode"] if tests_rows else "live"
    specialty = tests_rows[0]["specialty"] if tests_rows else "Лечебное дело"
    question_ids = (
        storage.client_entity_ids_bulk(
            source_mode,
            specialty,
            "test_question",
            [row["source_id"] for row in tests_rows],
            start_at=1,
        )
        if stable_ids
        else {row["source_id"]: index for index, row in enumerate(tests_rows, 1)}
    )
    prepared_tests = sorted(
        ((question_ids[row["source_id"]], row) for row in tests_rows),
        key=lambda item: item[0],
    )
    for question_id, row in prepared_tests:
        payload = row["payload"]
        ordered_options = _ordered_test_options(payload)
        tests_sheet.append([
            question_id,
            normalize_question_text(str(payload.get("question") or "")),
            *[option.get("text") or "" for _index, option in ordered_options],
        ])
    _style_tests_sheet(tests_sheet)


def _add_case_sheets(
    workbook: Workbook,
    storage: Storage,
    case_rows: List[Dict[str, Any]],
    stable_ids: bool = True,
) -> None:
    exercise_sheet = workbook.create_sheet("exercise_rows")
    questions_sheet = workbook.create_sheet("questions")
    answers_sheet = workbook.create_sheet("answers")
    exercise_sheet.append(["id", "header", "condition", "subject_id"])
    questions_sheet.append(["id", "question", "single_answer", "exercise_id", "text_before", "text_after", "header", "order"])
    answers_sheet.append(["id", "answer", "is_correct", "question_id"])

    source_mode = case_rows[0]["source_mode"] if case_rows else "live"
    specialty = case_rows[0]["specialty"] if case_rows else "Лечебное дело"
    exercise_ids = (
        storage.client_entity_ids_bulk(
            source_mode,
            specialty,
            "case_exercise",
            [row["source_id"] for row in case_rows],
            start_at=1,
        )
        if stable_ids
        else {row["source_id"]: index for index, row in enumerate(case_rows, 1)}
    )
    question_keys = []
    for row in case_rows:
        for order, question in enumerate(row["payload"].get("questions") or [], 1):
            question_key = "{}:question:{}".format(
                row["source_id"], question.get("question_number") or order
            )
            question_keys.append(question_key)
    question_ids = (
        storage.client_entity_ids_bulk(
            source_mode,
            specialty,
            "case_question",
            question_keys,
            start_at=600000,
        )
        if stable_ids
        else {key: 600000 + index for index, key in enumerate(dict.fromkeys(question_keys))}
    )
    prepared_cases = [(exercise_ids[row["source_id"]], row) for row in case_rows]
    for exercise_id, row in sorted(prepared_cases, key=lambda item: item[0]):
        payload = row["payload"]
        header = (
            extract_case_diagnosis(payload)
            or payload.get("header")
            or "Ситуационная задача {}".format(exercise_id)
        )
        exercise_sheet.append([exercise_id, header, _condition(payload), None])
        for order, question in enumerate(payload.get("questions") or [], 1):
            question_key = "{}:question:{}".format(
                row["source_id"], question.get("question_number") or order
            )
            current_question_id = question_ids[question_key]
            questions_sheet.append([
                current_question_id,
                normalize_question_text(str(question.get("question") or "")),
                not bool(question.get("multiple", False)),
                exercise_id,
                normalize_text(str(question.get("text_before") or "")) or None,
                _case_text_after(payload, question, order),
                None,
                order,
            ])
            for option_index, option in enumerate(question.get("options") or [], 1):
                answer_id = int("{}{}".format(current_question_id, option_index))
                answers_sheet.append([
                    answer_id,
                    option.get("text") or "",
                    option.get("is_correct") is True,
                    current_question_id,
                ])

    _style_template_case_sheet(exercise_sheet, {"A": 12.57, "B": 42, "C": 110, "D": 13})
    _style_template_case_sheet(questions_sheet, {"A": 12.57, "B": 33, "C": 21, "D": 12.57, "E": 8.71, "F": 60, "G": 13, "H": 10})
    _style_template_case_sheet(answers_sheet, {"A": 12.57, "B": 39, "C": 12.57, "D": 12.57})
    for row_number in range(2, questions_sheet.max_row + 1):
        text_after = str(questions_sheet.cell(row_number, 6).value or "")
        visual_lines = text_after.count("\n") + 1 + len(text_after) // 90
        questions_sheet.row_dimensions[row_number].height = min(300, max(45, visual_lines * 15))


def export_tests_xlsx(storage: Storage, run_id: str) -> Path:
    """Export tests as a separate workbook matching the client's test template."""
    ensure_directories()
    run, tests_rows, _case_rows = _ready_rows(storage, run_id)
    workbook = Workbook()
    _add_tests_sheet(workbook, storage, tests_rows, stable_ids=run.get("document_mode") != "new")
    target = EXPORT_DIR / "{}-tests.xlsx".format(run_id)
    workbook.save(target)
    return target


def export_cases_xlsx(storage: Storage, run_id: str) -> Path:
    """Export full situational cases as the client's three-sheet workbook."""
    ensure_directories()
    run, _tests_rows, case_rows = _ready_rows(storage, run_id)
    workbook = Workbook()
    workbook.remove(workbook.active)
    _add_case_sheets(workbook, storage, case_rows, stable_ids=run.get("document_mode") != "new")
    target = EXPORT_DIR / "{}-cases.xlsx".format(run_id)
    workbook.save(target)
    return target


def export_xlsx(storage: Storage, run_id: str) -> Path:
    """Backward-compatible combined workbook with both client template formats."""
    ensure_directories()
    run, tests_rows, case_rows = _ready_rows(storage, run_id)
    workbook = Workbook()
    stable_ids = run.get("document_mode") != "new"
    _add_tests_sheet(workbook, storage, tests_rows, stable_ids=stable_ids)
    _add_case_sheets(workbook, storage, case_rows, stable_ids=stable_ids)

    target = EXPORT_DIR / "{}.xlsx".format(run_id)
    workbook.save(target)
    return target


def export_images_zip(storage: Storage, run_id: str, image_dir: Path = IMAGE_DIR) -> Path:
    """Export downloaded case images under their exact client-facing filenames."""
    ensure_directories()
    _run, rows = _export_rows(storage, run_id)
    assets = _image_manifest(rows, Path(image_dir))
    target = EXPORT_DIR / "{}-images.zip".format(run_id)
    manifest = {
        "run_id": run_id,
        "image_count": len(assets),
        "directory": "images",
        "images": [
            {
                key: value
                for key, value in asset.items()
                if key not in {"source_path", "storage_name"}
            }
            for asset in assets
        ],
    }
    with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for asset in assets:
            archive.write(asset["source_path"], "images/{}".format(asset["filename"]))
        archive.writestr(
            "image_manifest.json",
            json.dumps(manifest, ensure_ascii=False, indent=2),
        )
    return target


def export_audit_xlsx(storage: Storage, run_id: str) -> Path:
    """Export a non-destructive catalog actuality report for one run."""
    ensure_directories()
    audit = storage.catalog_audit_items(
        run_id,
        kind="all",
        audit_status="all",
        limit=100000,
        offset=0,
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Актуальность"
    sheet.append([
        "Тип материала",
        "Исходный ID",
        "Статус проверки",
        "Последняя встреча",
        "Встречался, раз",
        "Версия",
        "Материал",
    ])
    kind_labels = {"test": "Тест", "case": "Ситуационная задача"}
    status_labels = {"confirmed": "Подтверждён в запуске", "not_checked": "Не проверен в запуске"}
    for item in audit["items"]:
        try:
            last_seen = datetime.fromisoformat(item["last_seen_at"]).astimezone().replace(tzinfo=None)
        except (TypeError, ValueError):
            last_seen = item["last_seen_at"]
        sheet.append([
            kind_labels.get(item["kind"], item["kind"]),
            item["source_id"],
            status_labels[item["audit_status"]],
            last_seen,
            item["times_seen"],
            item["version"],
            item["title"],
        ])

    header_fill = PatternFill("solid", fgColor="1F4E79")
    warning_fill = PatternFill("solid", fgColor="FFF2CC")
    confirmed_fill = PatternFill("solid", fgColor="E2EFDA")
    thin = Side(style="thin", color="D9D9D9")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Carlito", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in range(2, sheet.max_row + 1):
        status = sheet.cell(row, 3).value
        for column in range(1, 8):
            cell = sheet.cell(row, column)
            cell.font = Font(name="Carlito", size=11)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.fill = confirmed_fill if status == "Подтверждён в запуске" else warning_fill
        sheet.cell(row, 4).number_format = "yyyy-mm-dd hh:mm"
        sheet.row_dimensions[row].height = 32
    for column, width in {
        "A": 24, "B": 22, "C": 25, "D": 27, "E": 16, "F": 10, "G": 85,
    }.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    target = EXPORT_DIR / "{}-actuality.xlsx".format(run_id)
    workbook.save(target)
    return target
