"""Non-destructive comparison of a client workbook with one specialty bank."""
import html
import json
import logging
import re
import shutil
import threading
import time
import unicodedata
import uuid
import zipfile
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .exporter import _ordered_test_options
from .specialties import normalize_specialty

MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 100000
MAX_COLS = 100
RETENTION = 7 * 86400
ACTIVE = {"reading", "queued", "comparing", "exporting"}
LOG = logging.getLogger(__name__)


class ComparisonError(ValueError):
    pass


class _QuestionHTML(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []

    def handle_starttag(self, tag, attrs):
        if tag in {"sup", "sub"}:
            self.parts.append("^{" if tag == "sup" else "_{")
        elif tag in {"br", "p", "div", "li", "tr", "td"}:
            self.parts.append(" ")

    def handle_endtag(self, tag):
        if tag in {"sup", "sub"}:
            self.parts.append("}")
        elif tag in {"p", "div", "li", "tr", "td"}:
            self.parts.append(" ")

    def handle_data(self, data):
        self.parts.append(data)


def question_key(value):
    text = unicodedata.normalize("NFC", str(value or ""))
    parser = _QuestionHTML()
    parser.feed(html.unescape(text))
    parser.close()
    text = "".join(parser.parts)
    for chars, marker in (("⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾", "^"), ("₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎", "_")):
        table = str.maketrans(chars, "0123456789+-=()")
        text = re.sub("[" + chars + "]+", lambda m: marker + "{" + m[0].translate(table) + "}", text)
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = text.translate(str.maketrans({**dict.fromkeys("«»“”„", '"'),
                                       **dict.fromkeys("‘’", "'"),
                                       **dict.fromkeys("‐‑‒–—−", "-")}))
    return re.sub(r"\s+", " ", text).strip().casefold()


def read_client(path):
    """Only cell values; formulas are rejected, never executed or silently lost."""
    try:
        with zipfile.ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > 128 * 1024 * 1024:
                raise ComparisonError("Распакованный файл превышает 128 МБ.")
        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    except ComparisonError:
        raise
    except Exception as exc:
        raise ComparisonError("Не удалось открыть Excel. Загрузите исправный файл .xlsx.") from exc
    result = {}
    total_rows = 0
    try:
        if len(book.worksheets) > 20:
            raise ComparisonError("В файле допускается не более 20 листов.")
        for sheet in book:
            # Do not trust inflated or underreported worksheet dimensions.
            sheet.reset_dimensions()
            iterator = sheet.iter_rows()
            first = next(iterator, ())
            if len(first) > MAX_COLS:
                raise ComparisonError("В таблице допускается не более 100 колонок.")
            headers = [str(c.value or "").strip() for c in first]
            while headers and not headers[-1]:
                headers.pop()
            keys = [h.casefold() for h in headers]
            if "questiontext" not in keys:
                continue
            if any(c.data_type == "f" for c in first):
                raise ComparisonError("В заголовках есть формулы. Сохраните их как значения.")
            if not all(headers) or len(set(keys)) != len(keys):
                raise ComparisonError("Заголовки колонок должны быть заполнены и не повторяться.")
            qcol = keys.index("questiontext")
            records = []
            skipped = 0
            for line, cells in enumerate(iterator, 2):
                total_rows += 1
                if total_rows > MAX_ROWS:
                    raise ComparisonError("В файле допускается не более 100 000 строк данных.")
                if len(cells) > MAX_COLS:
                    raise ComparisonError("В таблице допускается не более 100 колонок.")
                if any(c.data_type == "f" for c in cells):
                    raise ComparisonError("В таблице есть формулы. Сохраните их как значения и загрузите файл повторно.")
                values = [c.value for c in cells]
                if any(v is not None for v in values[len(headers):]):
                    raise ComparisonError("У заполненных колонок отсутствуют заголовки.")
                values = (values + [None] * len(headers))[:len(headers)]
                if not question_key(values[qcol]):
                    skipped += 1
                    continue
                if not isinstance(values[qcol], str):
                    raise ComparisonError("В строке {} вопрос должен быть текстом.".format(line))
                records.append({"line": line, "values": values, "key": question_key(values[qcol])})
            result[sheet.title] = {"headers": headers, "records": records, "skipped": skipped}
    finally:
        book.close()
    if not result:
        raise ComparisonError("Не найдена колонка questiontext в первой строке листа.")
    return result


def test_identity(row):
    payload = row['payload']
    options = payload.get('options') or []
    return (question_key(payload['question']),
            tuple(sorted(question_key(o.get('text')) for o in options if o.get('is_correct') is True)),
            tuple(sorted(question_key(o.get('text')) for o in options if o.get('is_correct') is not True)))


def unique_tests(rows):
    seen = set()
    result = []
    for row in rows:
        identity = test_identity(row)
        if identity not in seen:
            seen.add(identity)
            result.append(row)
    return result


def compare_records(client, bank):
    old, new = defaultdict(list), defaultdict(list)
    for row in client["records"]:
        old[row["key"]].append(row)
    for row in bank:
        key = question_key(row["payload"].get("question"))
        if not key:
            raise ComparisonError("В банке найден тест без текста вопроса.")
        new[key].append(row)
    intersection = old.keys() & new.keys()
    return {
        "missing_client": unique_tests([r for k, group in new.items() if k not in old for r in group]),
        "missing_parser": [r for k, group in old.items() if k not in new for r in group],
        "old_groups": old, "new_groups": new,
        "summary": {
            "client_rows": len(client["records"]), "bank_rows": len(bank),
            "client_unique": len(old), "bank_unique": len(new),
            "matched_unique": len(intersection),
            "matched_client_rows": sum(len(old[k]) for k in intersection),
            "matched_bank_rows": sum(len(new[k]) for k in intersection),
            "missing_client_unique": len(new.keys() - old.keys()),
            "missing_parser_unique": len(old.keys() - new.keys()),
            "missing_client_rows": sum(len(v) for k, v in new.items() if k not in old),
            "missing_parser_rows": sum(len(v) for k, v in old.items() if k not in new),
            "client_duplicate_groups": sum(len(v) > 1 for v in old.values()),
            "bank_duplicate_groups": sum(len(v) > 1 for v in new.values()),
            "client_duplicate_extra_rows": sum(len(v)-1 for v in old.values()),
            "bank_duplicate_extra_rows": sum(len(v)-1 for v in new.values()),
            "skipped_rows": client["skipped"],
            "exported_missing_client": len(unique_tests([r for k,g in new.items() if k not in old for r in g])),
        },
    }


def bank_values(row):
    options = _ordered_test_options(row["payload"])
    return [row["source_id"], row["payload"]["question"],
            *[o.get("text", "") for _, o in options]]


SUMMARY_LABELS = {
    "exported_missing_client": "Нет у клиента: строк после удаления одинаковых вариантов",
    "client_rows": "Строк с вопросами в файле клиента", "bank_rows": "Тестов в снимке банка",
    "client_unique": "Уникальных текстов клиента", "bank_unique": "Уникальных текстов парсера",
    "matched_unique": "Совпадающих уникальных текстов", "matched_client_rows": "Совпавших строк клиента",
    "matched_bank_rows": "Совпавших тестов парсера", "missing_client_unique": "Нет у клиента: уникальных текстов",
    "missing_parser_unique": "Нет в парсере: уникальных текстов", "missing_client_rows": "Нет у клиента: строк",
    "missing_parser_rows": "Нет в парсере: строк", "client_duplicate_groups": "Групп дублей у клиента",
    "bank_duplicate_groups": "Групп дублей у парсера", "client_duplicate_extra_rows": "Повторных строк клиента",
    "bank_duplicate_extra_rows": "Повторных строк парсера", "skipped_rows": "Пропущено строк без вопроса",
}


def write_report(path, client, result, meta):
    book = Workbook()
    book.remove(book.active)

    def add(name, headers, rows):
        sheet = book.create_sheet(name)
        sheet.append(headers)
        for row_number, values in enumerate(rows, 2):
            sheet.append(values)
            # Force user-controlled strings to text, including leading '='.
            for cell in sheet.iter_rows(min_row=row_number, max_row=row_number, max_col=len(headers)):
                for item in cell:
                    if isinstance(item.value, str):
                        item.data_type = "s"
        for cell in sheet[1]:
            cell.data_type = "s"
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="087F5B")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial", size=11)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        from openpyxl.utils import get_column_letter
        for i, h in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(i)].width = 58 if ("question" in h or h in {"Вопрос", "Значение"}) else 26
        return sheet

    bh = ["source_id", "Вопрос", "Правильный ответ", "Ответ 2", "Ответ 3", "Ответ 4"]
    add("Нет в базе клиента", bh, (bank_values(r) for r in result["missing_client"]))
    add("Нет в банке парсера", client["headers"], (r["values"] for r in result["missing_parser"]))
    summary = [["Специальность", meta["specialty"]], ["Снимок банка (UTC)", meta["captured_at"]],
               ["Файл клиента", meta["filename"]], ["Лист клиента", meta["sheet"]],
               ["Правило", "Совпадение по нормализованному тексту вопроса; ответы не сравниваются."],
               ["Статус отсутствия", "Не найдено в текущем банке парсера не означает удаление из первоисточника."],
               ["Повторы парсера", "Одинаковый вопрос, правильный ответ и набор неверных ответов выводятся один раз. Порядок ответов не учитывается. Исходные повторы сохранены на листе Дубли."],
               ["Различия ответов", "Одинаковый текст с отличающимися ответами не объединяется. Такие группы также приведены на листе Дубли для проверки."],
               ["Версия сравнения", "2"]]
    summary.extend([SUMMARY_LABELS[k], v] for k, v in result["summary"].items())
    add("Сводка", ["Показатель", "Значение"], summary)
    duplicates = []
    for source, groups in (("Клиент", result["old_groups"]), ("Парсер", result["new_groups"])):
        for n, group in enumerate((g for g in groups.values() if len(g)>1), 1):
            for row in group:
                if source == "Клиент":
                    duplicates.append([source, n, len(group), row["line"], *([None]*6), *row["values"]])
                else:
                    duplicates.append([source, n, len(group), None, *bank_values(row), *([None]*len(client["headers"]))])
    add("Дубли", ["Источник", "Группа", "Строк в группе", "Строка Excel", *bh,
                   *["Клиент: " + h for h in client["headers"]]], duplicates)
    book.save(path)
    book.close()


class ComparisonManager:
    def __init__(self, storage, root):
        self.storage, self.root = storage, Path(root)
        self.root.mkdir(parents=True, exist_ok=True)
        self.lock = threading.RLock()
        self.pool = ThreadPoolExecutor(max_workers=1, thread_name_prefix="comparison")
        for directory in self.root.iterdir():
            if re.fullmatch(r"[0-9a-f]{32}", directory.name) and directory.is_dir():
                try:
                    meta = self.get(directory.name)
                    if meta["status"] in ACTIVE:
                        self.update(directory.name, status="interrupted", error="Сравнение прервано перезапуском. Загрузите файл повторно.")
                except (OSError, ValueError):
                    pass
        self.cleanup()

    def directory(self, identity):
        if not re.fullmatch(r"[0-9a-f]{32}", identity):
            raise ComparisonError("Отчёт не найден или срок хранения истёк.")
        return self.root / identity

    def get(self, identity):
        with self.lock:
            try:
                meta = json.loads((self.directory(identity) / "meta.json").read_text(encoding="utf-8"))
            except FileNotFoundError as exc:
                raise ComparisonError("Отчёт не найден или срок хранения истёк.") from exc
            if time.time() - meta["created"] > RETENTION and meta["status"] not in ACTIVE:
                raise ComparisonError("Срок хранения отчёта истёк. Загрузите файл повторно.")
            return meta

    def update(self, identity, **values):
        with self.lock:
            path = self.directory(identity)
            target = path / "meta.json"
            meta = json.loads(target.read_text(encoding="utf-8")) if target.exists() else {}
            meta.update(values)
            temp = path / "meta.tmp"
            temp.write_text(json.dumps(meta, ensure_ascii=False, default=str), encoding="utf-8")
            temp.replace(target)
            return meta

    def cleanup(self):
        for d in self.root.iterdir():
            if not re.fullmatch(r"[0-9a-f]{32}", d.name) or not d.is_dir():
                continue
            try:
                meta = json.loads((d / "meta.json").read_text(encoding="utf-8"))
                if time.time()-meta["created"] > RETENTION and meta["status"] not in ACTIVE:
                    shutil.rmtree(d)
            except (OSError, ValueError, KeyError):
                continue

    def upload(self, data, filename):
        if not filename.lower().endswith(".xlsx") or len(data) > MAX_BYTES or not data:
            raise ComparisonError("Выберите файл .xlsx размером до 20 МБ.")
        with self.lock:
            self.cleanup()
            self.assert_idle()
            identity = uuid.uuid4().hex
            d = self.directory(identity)
            d.mkdir()
            (d / "input.xlsx").write_bytes(data)
            meta = self.update(identity, id=identity, created=time.time(), status="reading",
                               filename=filename.replace("\\", "/").split("/")[-1][:200])
            self.pool.submit(self.inspect, identity)
            return meta

    def assert_idle(self):
        for d in self.root.glob("*/meta.json"):
            if not re.fullmatch(r"[0-9a-f]{32}", d.parent.name):
                continue
            if json.loads(d.read_text(encoding="utf-8")).get("status") in ACTIVE:
                raise ComparisonError("Другое сравнение ещё выполняется. Дождитесь его завершения.")

    def inspect(self, identity):
        try:
            sheets = read_client(self.directory(identity) / "input.xlsx")
            self.update(identity, status="ready", sheets=[{
                "name": name, "rows": len(s["records"]), "skipped": s["skipped"],
                "headers": s["headers"], "preview": [r["values"] for r in s["records"][:3]],
            } for name, s in sheets.items()])
        except Exception as exc:
            self.fail(identity, exc)

    def start(self, identity, sheet, specialty):
        specialty = normalize_specialty(specialty)
        with self.lock:
            meta = self.get(identity)
            if meta["status"] != "ready":
                raise ComparisonError("Для нового сравнения загрузите файл повторно.")
            if sheet not in [s["name"] for s in meta["sheets"]]:
                raise ComparisonError("Выберите лист из загруженного файла.")
            self.assert_idle()
            self.update(identity, status="queued", sheet=sheet, specialty=specialty)
            self.pool.submit(self.run, identity)
        return self.get(identity)

    def run(self, identity):
        try:
            meta = self.update(identity, status="comparing")
            snapshot = self.storage.comparison_snapshot(meta["specialty"])
            if not snapshot["items"]:
                raise ComparisonError("Банк выбранной специальности пуст. Сначала соберите тесты.")
            client = read_client(self.directory(identity) / "input.xlsx")[meta["sheet"]]
            if not client["records"]:
                raise ComparisonError("На выбранном листе нет заполненных вопросов.")
            result = compare_records(client, snapshot["items"])
            meta = self.update(identity, status="exporting", captured_at=snapshot["captured_at"], summary=result["summary"])
            write_report(self.directory(identity) / "report.xlsx", client, result, meta)
            self.update(identity, status="completed")
        except Exception as exc:
            self.fail(identity, exc)

    def fail(self, identity, exc):
        if not isinstance(exc, ComparisonError):
            LOG.exception("Comparison %s failed", identity)
        message = str(exc) if isinstance(exc, ComparisonError) else "Не удалось обработать файл. Проверьте таблицу и повторите загрузку."
        self.update(identity, status="failed", error=message)
